# -*- coding: utf-8 -*-
"""DOCX → 内核 Markdown（docs/plan/10 §5.1–5.2，单元 F3-docx-import-templates）。

import_docx(data, figs_dir) -> (content_md, ImportReport)：标题/段落/表格/图片/OMML 尽力映射为内核语法，
无法映射的原样保留并列入 report.items；三书按「摘要 / 权利要求书 / 说明书」标题启发式切分。
extract_styles(data) -> StyleSuggestion：正文/标题字体字号、页边距、行距 → settings.style 建议。
本模块只读 DOCX（python-docx + lxml），不写项目目录；写盘由 app.importer 负责。

映射（§5.1）：
  书标题        节页眉标签（内核自产：说明书摘要/摘要附图/权利要求书/说明书/说明书附图）或正文中
                独占一段的「摘要/说明书摘要/权利要求书/说明书」→ 切书；缺书 → book.missing，内容归说明书
  发明名称      Title/Heading 1（含中文 Word 的样式 id "1"）或首个居中标题字体段 → `# `；否则 title.guessed
  章节标题      Heading 2 → ###、Heading 3+ → ####；内核自产（全段标题字体、无样式）按段后距 8pt/6pt 分 ###/####
  段落          一段一行；去掉 [0001] 段号、零宽字符；下标 run → `_x` / `_{...}`
  权利要求      `^\\d+[.．、]` 起新权项（空行分隔），其余行并入当前权项
  表格          1×3 无框且中格 OMML → `$$…$$ 编号`；1×2 无框右格内嵌表 → 段号居左；其余 → Markdown 表
                （vMerge 续格 → ^^；gridSpan 补空格并报 table.merged；三线/重复表头/列宽 由 XML 还原）
  图片          写 figs_dir/图N.<ext>（同一图片部件只落一次）；附图/摘要附图节的说明文字不进正文
  行内公式      线性化进正文并报 formula.inline；OMML 含不可还原结构 → formula.partial（detail 为原 OMML）
  脚注/尾注/文本框/域/SmartArt → 丢弃并报 unsupported.<kind>（含段序号与原文片段）
"""
import io
import os
import re
import zipfile
from collections import Counter

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from lxml import etree
from pydantic import BaseModel

from patent_studio.kernel.content import table_spec_line
from patent_studio.kernel.inline import SUB_RE, SUP_RE
from patent_studio.kernel.style import DEFAULT_STYLE

UNIT = "F3-docx-import-templates"

BOOK_HEADINGS = {
    "abstract": ("摘要", "说明书摘要"),
    "claims": ("权利要求书",),
    "description": ("说明书",),
}
# 图片专用部分（内核 PART_KEYS 的另两项）：文字只作图注，不进正文
_FIG_PARTS = {"abstract_figure": ("摘要附图",), "figures": ("说明书附图", "附图")}
_HEADING_TO_BOOK = {h: b for b, hs in {**BOOK_HEADINGS, **_FIG_PARTS}.items() for h in hs}
_MD_HEADING = {"abstract": "## 摘要", "claims": "## 权利要求书", "description": "## 说明书"}

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_WP_EXTENT = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
_V_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"
_DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_TXBX = qn("w:txbxContent")


def _m(tag):
    return "{%s}%s" % (M_NS, tag)


_CONTAINERS = {qn(t) for t in ("w:hyperlink", "w:ins", "w:smartTag", "w:sdt", "w:sdtContent", "w:fldSimple",
                               "w:customXml", "w:moveTo")}
_HEADING_FONTS = {"黑体", "SimHei"}
_STRONG = set("+\u2212=\u2264\u2190\u2208\uff0c,")
_OPEN = {"(": ")", "[": "]", "{": "}"}
_CLOSE = {v: k for k, v in _OPEN.items()}
_ZERO_WIDTH = re.compile("[\u200b\u2060]")
_PARA_LABEL = re.compile(r"^\[\d{4}\][\u3000 \t]*")
_CLAIM_START = re.compile(r"^\d+\s*[.．、]")
# 与 builder/omml/inline 同一正则：反向生成时用它模拟正向解析，保证 `_x` / `^x` 形式能被原样识别
_SUB_RE = SUB_RE
_SCRIPT_RE = {"sub": ("_", SUB_RE), "sup": ("^", SUP_RE)}
# run 级 b/i/u → 内核记号（由外向内；b+i 合写 ***x***，inline._TOKEN_RE 专门识别）
_WRAP = {"b": ("**", "**"), "i": ("*", "*"), "u": ("<u>", "</u>")}
_NO_MARKS = frozenset()
_SUB_BASE = re.compile(r"[A-Za-z0-9]+|\S")
_SQRT_ARG = re.compile(r"[^\s/()\[\]{}√_+\u2212=\u2264\u2190\u2208\uff0c,]+(?:_[A-Za-z0-9,]+|_\{[^}]*\})?")
_FIG_CAPTION = re.compile(r"^图\s*\d+")
_STYLE_LEVEL = re.compile(r"^(?:Heading|heading)?\s*(\d)$")

EMU_PER_CM = 360000


class DocxImportError(Exception):
    """输入不是可读的 DOCX。"""


class ReportItem(BaseModel):
    code: str  # title.guessed | book.missing | formula.partial | formula.inline | unsupported.<kind>
    para_index: int
    text: str
    detail: str | None = None


class StyleSuggestion(BaseModel):
    style: dict = {}  # 与 patent_studio.kernel.style.DEFAULT_STYLE 同形，只含检测到的键
    evidence: dict[str, str] = {}


class ImportReport(BaseModel):
    books: dict[str, int] = {}  # 各书行数
    figures: list[str] = []
    items: list[ReportItem] = []
    style_suggestion: StyleSuggestion | None = None
    figure_settings: dict = {}  # {abstract_figure: [名, 宽cm], spec_figures: [[名, 宽cm]…]}，供 importer 写 studio.json
    sub_except: list[str] = []  # 正文中字面出现、不应被内核解成下标的标识符（如 tsf_guard）


def _open(data):
    try:
        return Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as e:
        raise DocxImportError("不是有效的 DOCX 文件") from e


# ---------------- 通用 XML 读取 ----------------

def _val(el, default=None):
    return default if el is None else el.get(qn("w:val"), default)


def _blocks(container):
    """容器（body / tc / sdtContent）的块级子元素：w:p 与 w:tbl，穿透 w:sdt。"""
    for el in container:
        if el.tag in (qn("w:p"), qn("w:tbl")):
            yield el
        elif el.tag == qn("w:sdt"):
            content = el.find(qn("w:sdtContent"))
            if content is not None:
                yield from _blocks(content)


def _border_vals(pr, tag):
    b = pr.find(qn(tag)) if pr is not None else None
    if b is None:
        return None
    return {edge.tag.split("}")[1]: _val(edge, "single") for edge in b}


def _borderless(tbl):
    b = _border_vals(tbl.find(qn("w:tblPr")), "w:tblBorders")
    return b is not None and all(v in ("none", "nil") for v in b.values())


def _wtext(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _snippet(text, n=40):
    text = re.sub(r"\s+", " ", text).strip()
    return text if len(text) <= n else text[:n] + "…"


class _Run:
    """一个 w:r 的文本与格式（供正文拼接、标题判定与样式统计共用）。"""

    __slots__ = ("text", "sub", "sup", "ea", "ascii", "size", "bold", "italic", "underline")

    def __init__(self, text, rpr):
        self.text = text
        va = _val(rpr.find(qn("w:vertAlign"))) if rpr is not None else None
        self.sub = va == "subscript"
        self.sup = va == "superscript"
        self.italic = _on(rpr.find(qn("w:i")) if rpr is not None else None)
        u = rpr.find(qn("w:u")) if rpr is not None else None
        self.underline = u is not None and u.get(qn("w:val"), "single") != "none"
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        self.ea = rf.get(qn("w:eastAsia")) if rf is not None else None
        self.ascii = rf.get(qn("w:ascii")) if rf is not None else None
        sz = _val(rpr.find(qn("w:sz"))) if rpr is not None else None
        self.size = _num(int(sz) / 2, 1) if sz and sz.isdigit() else None
        self.bold = _on(rpr.find(qn("w:b")) if rpr is not None else None)

    @property
    def marks(self):
        """选区级格式（b/i/u）→ inline.MARKS 子集；下标/上标单独走 pieces 的 kind。"""
        return frozenset(m for m, on in (("b", self.bold), ("i", self.italic), ("u", self.underline)) if on)


def _on(toggle):
    """w:b / w:i 等开关属性：存在且 val 不为 0/false 即开。"""
    return toggle is not None and toggle.get(qn("w:val"), "1") not in ("0", "false")


class _Para:
    """段落读取结果：pieces 为 ('t'|'sub'|'sup'|'math', 文本) 序列（'\\n' = 软回车），piece_marks 为与之对齐的
    run 级 b/i/u 标志集（math 片段为空集），images 为 (rId, cx) 列表。"""

    def __init__(self, p):
        self.el = p
        ppr = p.find(qn("w:pPr"))
        self.style_id = _val(ppr.find(qn("w:pStyle"))) if ppr is not None else None
        self.jc = _val(ppr.find(qn("w:jc"))) if ppr is not None else None
        sp = ppr.find(qn("w:spacing")) if ppr is not None else None
        self.after = int(sp.get(qn("w:after"))) if sp is not None and sp.get(qn("w:after")) else None
        self.line = sp.get(qn("w:line")) if sp is not None else None
        self.line_rule = sp.get(qn("w:lineRule")) if sp is not None else None
        ind = ppr.find(qn("w:ind")) if ppr is not None else None
        self.first_line = int(ind.get(qn("w:firstLine"))) if ind is not None and ind.get(qn("w:firstLine")) else None
        self.has_sect_break = ppr is not None and ppr.find(qn("w:sectPr")) is not None
        self.runs: list[_Run] = []
        self.pieces: list[tuple[str, str]] = []
        self.piece_marks: list[frozenset] = []
        self.images: list[tuple[str, int | None]] = []
        self.unsupported: list[tuple[str, str]] = []  # (kind, detail)
        self.inline_math: list[etree._Element] = []

    @property
    def text(self):
        return "".join(t for _, t in self.pieces)

    @property
    def text_runs(self):
        return [r for r in self.runs if r.text.strip()]

    def all_heading_font(self):
        rs = self.text_runs
        return bool(rs) and all(r.ea in _HEADING_FONTS for r in rs)

    def level(self):
        """标题层级：0 = 发明名称级，1 = ###，2 = ####；None = 非标题。"""
        if self.style_id:
            if self.style_id.lower() == "title":
                return 0
            m = _STYLE_LEVEL.match(self.style_id)
            if m:
                return min(int(m.group(1)) - 1, 2)
        if self.all_heading_font() and len(self.text) <= 60 and "。" not in self.text:
            if self.jc == "center":
                return 0
            return 2 if self.after == 120 else 1
        return None


def _read_para(p, linear_omml):
    """读段落（也用于表格单元格）。linear_omml(oMath) -> str 由调用方注入（需要记录报告）。"""
    para = _Para(p)

    def walk(el):
        for ch in el:
            tag = ch.tag
            if tag == qn("w:r"):
                run(ch)
            elif tag in _CONTAINERS:
                if tag == qn("w:fldSimple"):
                    para.unsupported.append(("field", ch.get(qn("w:instr")) or ""))
                walk(ch)
            elif tag == _m("oMath"):
                para.inline_math.append(ch)
                piece("math", linear_omml(ch), _NO_MARKS)
            elif tag == _m("oMathPara"):
                for om in ch:
                    if om.tag == _m("oMath"):
                        para.inline_math.append(om)
                        piece("math", linear_omml(om), _NO_MARKS)

    def piece(kind, text, marks):
        para.pieces.append((kind, text))
        para.piece_marks.append(marks)

    def run(r):
        rpr = r.find(qn("w:rPr"))
        buf = []
        for ch in r:
            tag = ch.tag
            if tag == qn("w:t"):
                buf.append(ch.text or "")
            elif tag == qn("w:tab"):
                buf.append(" ")
            elif tag == qn("w:br"):
                if ch.get(qn("w:type")) in (None, "textWrapping"):
                    buf.append("\n")
            elif tag == qn("w:noBreakHyphen"):
                buf.append("-")
            elif tag == qn("w:drawing"):
                drawing(ch)
            elif tag == qn("w:pict"):
                img = ch.find(".//" + _V_IMAGEDATA)
                if img is not None and img.get(_R_ID):
                    para.images.append((img.get(_R_ID), None))
                elif ch.find(".//" + _TXBX) is not None:
                    para.unsupported.append(("textbox", _wtext(ch)))
                else:
                    para.unsupported.append(("pict", ""))
            elif tag == qn("w:footnoteReference"):
                para.unsupported.append(("footnote", ch.get(qn("w:id")) or ""))
            elif tag == qn("w:endnoteReference"):
                para.unsupported.append(("endnote", ch.get(qn("w:id")) or ""))
            elif tag == qn("w:instrText"):
                para.unsupported.append(("field", ch.text or ""))
        text = "".join(buf)
        rn = _Run(text, rpr)
        para.runs.append(rn)
        if text:
            piece("sub" if rn.sub else "sup" if rn.sup else "t", text, rn.marks)

    def drawing(d):
        if d.find(".//" + _TXBX) is not None:
            para.unsupported.append(("textbox", _wtext(d)))
            return
        if d.find(".//{%s}relIds" % _DGM_NS) is not None:
            para.unsupported.append(("smartart", ""))
            return
        blip = d.find(".//" + _A_BLIP)
        if blip is None or not blip.get(_R_EMBED):
            para.unsupported.append(("drawing", ""))
            return
        ext = d.find(".//" + _WP_EXTENT)
        cx = int(ext.get("cx")) if ext is not None and ext.get("cx") else None
        para.images.append((blip.get(_R_EMBED), cx))

    walk(p)
    return para


def _join_pieces(pieces, marks=None):
    """('t'|'sub'|'sup', 文本) 序列 → 含内核行内记号的文本。
    下标/上标优先写成 `_x` / `^x`；只有当正向解析会多吃/少吃后文时才加花括号 `_{x}` / `^{x}`。
    marks（与 pieces 对齐的 b/i/u 标志集，缺省全空）：相邻同标志片段合为一个跨度，包上 `**` / `*` / `<u>`，
    跨度两端的空白留在记号外（`*` 两侧不得是空白）；软回车 '\\n' 处断开（标记不跨行）。"""
    if marks is None:
        marks = [_NO_MARKS] * len(pieces)
    flat = []  # (kind, text, marks)，空片段去掉，空白/换行从标志片段里剥出
    for (kind, t), mk in zip(pieces, marks):
        if not t:
            continue
        if not mk or kind == "math":
            flat.append((kind, t, _NO_MARKS))
            continue
        for part in re.split(r"(\n)", t):
            if part:
                flat.append((kind, part, _NO_MARKS if part == "\n" else mk))
    out, i = "", 0
    while i < len(flat):
        mk = flat[i][2]
        j = i
        while j < len(flat) and flat[j][2] == mk:
            j += 1
        span = ""
        for k in range(i, j):
            kind, t, _ = flat[k]
            if kind not in _SCRIPT_RE:
                span += t
                continue
            sigil, rx = _SCRIPT_RE[kind]
            rest = "".join((_SCRIPT_RE[x[0]][0] if x[0] in _SCRIPT_RE else "") + x[1]
                           for x in flat[k + 1:j])
            m = rx.match(sigil + t + rest)
            span += sigil + t if m and m.group(1) == t else sigil + "{" + t + "}"
        core = span.strip()
        if mk and core:
            for m in ("u", "i", "b"):
                if m in mk:
                    core = _WRAP[m][0] + core + _WRAP[m][1]
            lead = span[:len(span) - len(span.lstrip())]
            span = lead + core + span[len(span.rstrip()):]
        out += span
        i = j
    return out


def _clean(text):
    return _PARA_LABEL.sub("", _ZERO_WIDTH.sub("", text)).strip()


# ---------------- OMML → 线性公式（omml.append_omml 的尽力逆变换） ----------------

def _mval(el, default):
    return default if el is None else el.get(_m("val"), default)


def _has_top_op(pieces, slash=True):
    """顶层（括号外、下标外）是否含边界运算符（slash 时也算斜线）。

    正向解析时分子/分母在二者处截断，求和体只在运算符处截断。
    """
    depth = 0
    for kind, s in pieces:
        if kind == "sub":
            continue
        for ch in s:
            if ch in _OPEN:
                depth += 1
            elif ch in _CLOSE:
                depth -= 1
            elif depth == 0 and (ch in _STRONG or (slash and ch == "/")):
                return True
    return False


class _Omml:
    def __init__(self):
        self.partial = False

    def box(self, el):
        return _join_pieces(self.pieces(el)).strip() if el is not None else ""

    def paren(self, el, slash=True):
        """参数位文本；含顶层运算符时加圆括号（正向 _unwrap/_strip_edges 会再去掉）。"""
        ps = self.pieces(el) if el is not None else []
        s = _join_pieces(ps).strip()
        return "(" + s + ")" if _has_top_op(ps, slash) else s

    def pieces(self, el):
        out = []
        for ch in el:
            tag = ch.tag
            if not tag.startswith("{%s}" % M_NS) or tag.endswith("Pr"):
                continue
            kind = tag.split("}")[1]
            if kind == "r":
                out.append(("t", "".join(t.text or "" for t in ch.iter(_m("t"))).replace("\u2223", "|")))
            elif kind == "f":
                out.append(("t", self.paren(ch.find(_m("num"))) + "/" + self.paren(ch.find(_m("den")))))
            elif kind == "sSub":
                base = self.pieces(ch.find(_m("e")))
                if not _SUB_BASE.fullmatch(_join_pieces(base).strip()):
                    self.partial = True
                out.extend(base)
                out.append(("sub", self.box(ch.find(_m("sub")))))
            elif kind == "rad":
                if self.box(ch.find(_m("deg"))):
                    self.partial = True
                arg = self.box(ch.find(_m("e")))
                out.append(("t", "√" + (arg if _SQRT_ARG.fullmatch(arg) else "(" + arg + ")")))
            elif kind == "nary":
                pr = ch.find(_m("naryPr"))
                chr_ = _mval(pr.find(_m("chr")) if pr is not None else None, "∫")
                if chr_ not in ("Σ", "∑") or self.box(ch.find(_m("sup"))):
                    self.partial = True
                body = self.paren(ch.find(_m("e")), slash=False)
                out.append(("t", "Σ_{%s}%s" % (self.box(ch.find(_m("sub"))), body)))
            elif kind == "d":
                pr = ch.find(_m("dPr"))
                beg = _mval(pr.find(_m("begChr")) if pr is not None else None, "(")
                end = _mval(pr.find(_m("endChr")) if pr is not None else None, ")")
                es = ch.findall(_m("e"))
                if len(es) != 1 or _OPEN.get(beg) != end:
                    self.partial = True
                out.append(("t", beg + "，".join(self.box(e) for e in es) + end))
            elif kind == "sSup":
                self.partial = True
                out.append(("t", self.box(ch.find(_m("e"))) + "^" + self.box(ch.find(_m("sup")))))
            elif kind == "sSubSup":
                base = self.pieces(ch.find(_m("e")))
                if not _SUB_BASE.fullmatch(_join_pieces(base).strip()):
                    self.partial = True
                out.extend(base)
                out.append(("sub", self.box(ch.find(_m("sub")))))
                out.append(("sup", self.box(ch.find(_m("sup")))))
            else:  # bar / acc / box / eqArr / func / groupChr / limLow / limUpp / m / sPre …
                self.partial = True
                out.append(("t", "".join(t.text or "" for t in ch.iter(_m("t")))))
        return out


# ---------------- 文档 → content.md ----------------

class _Importer:
    def __init__(self, doc, figs_dir):
        self.doc = doc
        self.figs_dir = figs_dir
        self.items: list[ReportItem] = []
        self.title = None
        self.book = None  # abstract | claims | description | abstract_figure | figures | None
        self.lines = {"abstract": [], "claims": [], "description": []}
        self.seen = set()
        self.fig_names: dict[str, str] = {}  # partname → 图N.ext
        self.figures: list[str] = []
        self.abstract_fig = None
        self.spec_figs: list[list] = []
        self.inline_figs: list[list] = []
        self.blk = 0
        self.sub_except: list[str] = []
        labels = {v: k for k, v in DEFAULT_STYLE["header"]["labels"].items()}
        self.parts = {}
        for i, sec in enumerate(doc.sections):
            text = "".join(p.text for p in sec.header.paragraphs).strip()
            if text in labels:
                self.parts[i] = labels[text]

    def report(self, code, text, detail=None):
        self.items.append(ReportItem(code=code, para_index=self.blk, text=_snippet(text), detail=detail))

    def literal_underscores(self, pieces):
        """非下标 run 中字面的 `x_y`：内核会把 `_y` 排成下标，需进 sub_except 才能原样重现。
        （字面 `*` / `^` / `<u>` 没有白名单机制，导入后会被内核当作行内记号；语料中不出现，见 docs/plan/22 §2。）"""
        for kind, t in pieces:
            if kind != "t":
                continue
            for m in _SUB_RE.finditer(t):
                base = re.search(r"[A-Za-z0-9]*$", t[:m.start()]).group(0)
                ident = base + m.group(0)
                if ident not in self.sub_except:
                    self.sub_except.append(ident)

    def linear(self, om):
        conv = _Omml()
        s = conv.box(om)
        if conv.partial:
            self.report("formula.partial", s, etree.tostring(om, encoding="unicode"))
        return s

    def run(self):
        si = 0
        if 0 in self.parts:
            self.enter(self.parts[0])
        for el in _blocks(self.doc.element.body):
            self.blk += 1
            if el.tag == qn("w:tbl"):
                self.table(el)
                continue
            para = _read_para(el, self.linear)
            self.paragraph(para)
            if para.has_sect_break:
                si += 1
                if si in self.parts:
                    self.enter(self.parts[si])
        return self.finish()

    def enter(self, book):
        self.book = book
        self.seen.add(book)

    # --- 段落 ---
    def paragraph(self, para):
        for kind, detail in para.unsupported:
            self.report("unsupported." + kind, detail or para.text)
        for rid, cx in para.images:
            self.image(rid, cx)
        if para.inline_math and self.book not in ("abstract_figure", "figures"):
            for om in para.inline_math:
                self.report("formula.inline", _Omml().box(om))
        text = _clean(_join_pieces(para.pieces, para.piece_marks))
        if not text:
            return
        compact = re.sub(r"\s+", "", text)
        if compact in _HEADING_TO_BOOK:
            self.enter(_HEADING_TO_BOOK[compact])
            return
        level = para.level()
        if self.title is None and level == 0 and self.book in (None, "abstract", "description"):
            self.title = text
            return
        if text == self.title and para.jc == "center":
            return
        if self.book in ("abstract_figure", "figures"):
            if not _FIG_CAPTION.match(text):
                self.report("unsupported.figures_text", text)
            return
        book = self.book or "description"
        lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
        if book == "description" and level is not None:
            self.lines[book].extend(["", ("### " if level <= 1 else "#### ") + " ".join(lines)])
            return
        if book == "description" and self.is_code(para):
            self.lines[book].extend(["", "```", *lines, "```"])
            return
        self.literal_underscores(para.pieces)
        self.emit(book, lines)

    def emit(self, book, lines):
        out = self.lines[book]
        if book == "abstract":
            if not out:
                out.append("")
            out.extend(lines)
        elif book == "claims":
            for ln in lines:
                if _CLAIM_START.match(ln) or not out:
                    out.append("")
                out.append(ln)
        else:
            for ln in lines:
                out.extend(["", ln])

    @staticmethod
    def is_code(para):
        rs = [r for r in para.text_runs if not _PARA_LABEL.match(r.text)]
        mono = DEFAULT_STYLE["fonts"]["mono"]
        return bool(rs) and all(r.ascii == mono for r in rs)

    # --- 图片 ---
    def image(self, rid, cx):
        try:
            part = self.doc.part.rels[rid].target_part
        except (KeyError, ValueError):
            self.report("unsupported.image", rid)
            return
        key = str(part.partname)
        if key not in self.fig_names:
            ext = os.path.splitext(key)[1] or ".bin"
            name = f"图{len(self.fig_names) + 1}{ext}"
            os.makedirs(self.figs_dir, exist_ok=True)
            with open(os.path.join(self.figs_dir, name), "wb") as f:
                f.write(part.blob)
            self.fig_names[key] = name
            self.figures.append(name)
        entry = [self.fig_names[key], round(cx / EMU_PER_CM, 2) if cx else 15.0]
        if self.book == "abstract_figure":
            self.abstract_fig = entry
        elif self.book == "figures":
            self.spec_figs.append(entry)
        else:
            self.inline_figs.append(entry)
            self.report("figure.inline", entry[0], "图片位于正文中，已存入 figs/，未在正文生成占位")

    # --- 表格 ---
    def table(self, tbl):
        rows = tbl.findall(qn("w:tr"))
        cells0 = rows[0].findall(qn("w:tc")) if rows else []
        if len(rows) == 1 and len(cells0) == 3 and cells0[1].find(".//" + _m("oMath")) is not None:
            self.equation(cells0[1].find(".//" + _m("oMath")), self.cell_text(cells0[2]))
            return
        inner = cells0[1].find(qn("w:tbl")) if len(rows) == 1 and len(cells0) == 2 else None
        if inner is not None and _borderless(tbl):
            self.data_table(inner, left_num=True)
            return
        self.data_table(tbl, left_num=False)

    def cell_text(self, tc, header=False):
        """header=True：表头行，内核 add_table 本就把首行排成粗体，导入时丢掉 run 级 bold 免得写出冗余 `**`。"""
        parts = []
        for blk in _blocks(tc):
            if blk.tag == qn("w:p"):
                para = _read_para(blk, self.linear)
                for kind, detail in para.unsupported:
                    self.report("unsupported." + kind, detail or para.text)
                for rid, cx in para.images:
                    self.image(rid, cx)
                self.literal_underscores(para.pieces)
                marks = [mk - {"b"} for mk in para.piece_marks] if header else para.piece_marks
                parts.append(_clean(_join_pieces(para.pieces, marks)))
            else:
                self.report("unsupported.nested_table", "")
        return re.sub(r"\s+", " ", " ".join(p for p in parts if p)).replace("|", "\u2223")

    def equation(self, om, num):
        formula = self.linear(om)
        book = self.book or "description"
        if book == "description":
            self.lines[book].extend(["", ("$$%s$$ %s" % (formula, num)).rstrip()])
        elif book == "claims":
            if num:
                self.report("unsupported.equation_number", num, "权利要求书公式不带编号")
            if not self.lines[book]:
                self.lines[book].append("")
            self.lines[book].append("$$%s$$" % formula)
        elif book == "abstract":
            self.report("formula.inline", formula, "摘要不支持独立公式，已线性化为正文")
            self.emit(book, [formula])

    def data_table(self, tbl, left_num):
        rows = tbl.findall(qn("w:tr"))
        grid, merged = [], False
        for ri, tr in enumerate(rows):
            cells = []
            for tc in tr.findall(qn("w:tc")):
                tcpr = tc.find(qn("w:tcPr"))
                vm = tcpr.find(qn("w:vMerge")) if tcpr is not None else None
                span = int(_val(tcpr.find(qn("w:gridSpan")) if tcpr is not None else None, "1"))
                if vm is not None and vm.get(qn("w:val"), "continue") == "continue":
                    cells.append("^^")
                else:
                    cells.append(self.cell_text(tc, header=ri == 0))
                if span > 1:
                    cells.extend([""] * (span - 1))
                    merged = True
            grid.append(cells)
        if not grid:
            return
        if merged:
            self.report("table.merged", grid[0][0] if grid[0] else "", "含横向合并单元格（gridSpan），内核以空格补齐")
        opts = []
        b = _border_vals(tbl.find(qn("w:tblPr")), "w:tblBorders")
        if b and b.get("top") not in (None, "none") and b.get("insideH") == "none" and b.get("left") == "none":
            opts.append("三线")
        trpr = rows[0].find(qn("w:trPr"))
        if trpr is not None and trpr.find(qn("w:tblHeader")) is not None:
            opts.append("重复表头")
        if left_num:
            opts.append("段号居左")
        widths = []
        for tc in rows[0].findall(qn("w:tc")):
            tcpr = tc.find(qn("w:tcPr"))
            w = tcpr.find(qn("w:tcW")) if tcpr is not None else None
            if w is None or w.get(qn("w:type")) != "pct":
                widths = None
                break
            widths.append(int(w.get(qn("w:w"))) / 50)
        if widths and not (left_num and len(set(widths)) == 1):
            opts.append("列宽=" + ",".join("%g" % x for x in widths))
        book = self.book or "description"
        lines = ["| " + " | ".join(r) + " |" for r in grid]
        if book == "description":
            self.lines[book].append("")
            if opts:
                self.lines[book].append(table_spec_line(opts))
            self.lines[book].extend(lines)
        else:
            self.report("unsupported.table", grid[0][0] if grid[0] else "", "表格只支持出现在说明书中，已按文本行保留")
            self.emit(book, lines)

    # --- 汇总 ---
    def finish(self):
        if self.title is None:
            self.title = "未命名发明"
            self.items.insert(0, ReportItem(code="title.guessed", para_index=0, text="", detail="未识别到发明名称"))
        for book in ("abstract", "claims", "description"):
            if book not in self.seen and not self.lines[book]:
                self.report("book.missing", _MD_HEADING[book][3:], "未找到该书标题，内容已归入说明书")
        out = ["# " + self.title, ""]
        for book in ("abstract", "claims", "description"):
            out.append(_MD_HEADING[book])
            out.extend(self.lines[book])
            out.append("")
        md = "\n".join(out).rstrip("\n") + "\n"
        spec = self.spec_figs or self.inline_figs
        fig_settings = {}
        if self.figures:
            fig_settings["abstract_figure"] = self.abstract_fig or (spec[0] if spec else [self.figures[0], 15.0])
            fig_settings["spec_figures"] = spec or [fig_settings["abstract_figure"]]
        report = ImportReport(
            books={b: sum(1 for ln in self.lines[b] if ln) for b in self.lines},
            figures=self.figures, items=self.items, figure_settings=fig_settings, sub_except=self.sub_except,
        )
        return md, report


def import_docx(data: bytes, figs_dir: str) -> tuple[str, ImportReport]:
    """图片写入 figs_dir（图N.<ext>）；返回 content.md 文本与报告（含 style_suggestion）。"""
    doc = _open(data)
    md, report = _Importer(doc, figs_dir).run()
    report.style_suggestion = _styles(doc)
    return md, report


# ---------------- 样式提取（§5.2） ----------------

def _num(x, nd):
    v = round(x, nd)
    return int(v) if v == int(v) else v


def _majority(counter):
    return counter.most_common(1)[0] if counter else None


def _styles(doc):
    body_ea, body_ascii, body_sz = Counter(), Counter(), Counter()
    head_ea, head_sz, title_sz = Counter(), Counter(), Counter()
    line, after, first = Counter(), Counter(), Counter()
    for el in _blocks(doc.element.body):
        if el.tag != qn("w:p"):
            continue
        para = _read_para(el, lambda om: "")
        runs = [r for r in para.text_runs if not _PARA_LABEL.match(r.text)]
        if not runs:
            continue
        level = para.level()
        if level == 0:
            for r in runs:
                if r.size:
                    title_sz[r.size] += len(r.text)
            continue
        if level is not None:
            for r in runs:
                if r.ea:
                    head_ea[r.ea] += len(r.text)
                if r.size:
                    head_sz[r.size] += len(r.text)
            continue
        for r in runs:
            n = len(r.text)
            if r.ea:
                body_ea[r.ea] += n
            if r.ascii:
                body_ascii[r.ascii] += n
            if r.size:
                body_sz[r.size] += n
        if para.line and para.line_rule in (None, "auto"):
            line[_num(int(para.line) / 240, 2)] += 1
        if para.after is not None:
            after[_num(para.after / 20, 1)] += 1
        if para.first_line:
            first[_num(para.first_line / 20, 1)] += 1
    if not body_ea and not body_ascii:  # 未内联字体时退到 Normal 样式的显式设置
        normal = doc.styles.element.find(".//" + qn("w:style") + "[@" + qn("w:styleId") + "='Normal']")
        rpr = normal.find(qn("w:rPr")) if normal is not None else None
        if rpr is not None:
            r = _Run("", rpr)
            if r.ea:
                body_ea[r.ea] += 1
            if r.ascii:
                body_ascii[r.ascii] += 1
            if r.size:
                body_sz[r.size] += 1

    style, ev = {}, {}

    def put(path, value, evidence):
        d = style
        for k in path[:-1]:
            d = d.setdefault(k, {})
        d[path[-1]] = value
        ev[".".join(path)] = evidence

    for path, counter, unit in ((("fonts", "body_cjk"), body_ea, "字"), (("fonts", "ascii"), body_ascii, "字"),
                                (("sizes", "body"), body_sz, "字"), (("fonts", "heading_cjk"), head_ea, "字"),
                                (("sizes", "section_heading"), head_sz, "字"), (("sizes", "title"), title_sz, "字"),
                                (("paragraph", "line_spacing"), line, "段"),
                                (("paragraph", "space_after", "body"), after, "段"),
                                (("paragraph", "first_indent_pt"), first, "段")):
        m = _majority(counter)
        if m:
            put(path, m[0], f"多数 run/段：{m[0]}（{m[1]} {unit}，共 {sum(counter.values())} {unit}）")
    if doc.sections:
        sec = doc.sections[0]
        if sec.page_width and sec.page_height:
            put(("page", "width_cm"), _num(sec.page_width.cm, 1), "首节 pgSz")
            put(("page", "height_cm"), _num(sec.page_height.cm, 1), "首节 pgSz")
        for key, val in (("margin_top_mm", sec.top_margin), ("margin_bottom_mm", sec.bottom_margin),
                         ("margin_left_mm", sec.left_margin), ("margin_right_mm", sec.right_margin)):
            if val is not None:
                put(("page", key), _num(val.mm, 1), "首节 pgMar")
    return StyleSuggestion(style=style, evidence=ev)


def extract_styles(data: bytes) -> StyleSuggestion:
    return _styles(_open(data))
