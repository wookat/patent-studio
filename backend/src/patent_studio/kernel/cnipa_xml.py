# -*- coding: utf-8 -*-
"""content.md → CNIPA 电子申请 XML（五书）导出。

目标格式来自国知局「WORD 转 XML 编辑器」安装包自带的官方 DTD
（cnipa_dtd/cn-application-body-20080416.dtd，ST.36 派生），输出五个独立 XML：

  100001.xml 权利要求书  cn-claims / claim[@num] / claim-text
  100002.xml 说明书      description / invention-title + 五部分 section / p[@num]
  100003.xml 说明书附图  cn-drawings / figure[@num] / img
  100004.xml 说明书摘要  cn-abstract / p[@num]
  100005.xml 摘要附图    cn-abstract / cn-abst-figure / figure / img

段号、权项号、图号与 builder.py 生成的 DOCX 完全同源（同一份 load_md() 结果、同一套计数
规则），因此 DOCX/PDF 预览里看到的 [0001] 就是 XML 里的 p/@num。

公式、表格、附图按官方插件的默认做法统一转为 300 dpi JPG：把所有此类对象各占一页放进一份
临时 DOCX（复用 PatentBuilder 的排版），LibreOffice 转 PDF，pdftoppm 栅格化后按内容裁边。
调用方需提供 docx_to_pdf(docx_path, out_dir) 回调（app.lo.docx_to_pdf）。

导出后用官方 DTD 逐书校验；DTD 自身 cn-drawings 内容模型在 libxml 下非确定性的报错
（"Content model of cn-drawings is not deterministic"）不计为错误。
"""
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from lxml import etree
from PIL import Image, ImageChops

from .builder import DEFAULT_SUB_EXCEPT, FigureError, PatentBuilder
from .content import load_md
from .inline import MARKS, inline_tokens
from .omml import latex_to_text
from .style import merge_style

_INLINE_MATH = re.compile(r"\$\$([^$\n]+?)\$\$|\$([^$\n]+?)\$")

DTD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cnipa_dtd")
DTD_PATH = os.path.join(DTD_DIR, "cn-application-body-20080416.dtd")
XSL_PATH = os.path.join(DTD_DIR, "showxml.xsl")

DOCTYPE = '<!DOCTYPE cn-application-body SYSTEM "/dtdandxsl/cn-application-body-20080416.dtd">'
STYLESHEET_PI = '<?xml-stylesheet type="text/xsl" href="/dtdandxsl/showxml.xsl"?>'

BOOKS = (("100001", "权利要求书"), ("100002", "说明书"), ("100003", "说明书附图"),
         ("100004", "说明书摘要"), ("100005", "摘要附图"))

# 说明书 ### 章节 → DTD section 元素；不在表内的章节退化为 heading + p 直接挂在 description 下
SECTION_ELEMENTS = {"技术领域": "technical-field", "背景技术": "background-art",
                    "发明内容": "disclosure", "附图说明": "description-of-drawings",
                    "具体实施方式": "mode-for-invention"}
NONDETERMINISTIC_DTD_MSG = "not deterministic"

DPI = 300
ASSET_KIND_NAMES = {"fig": "附图", "eq": "公式", "tbl": "表格"}
_CLAIM_REF_RE = re.compile(r"权利要求(\d+)(?:至(\d+))?")


class XmlExportError(Exception):
    """XML 导出阶段错误（渲染对象跨页、栅格化工具缺失等）。"""


class XmlContentError(XmlExportError):
    """内容满足 DOCX 排版但不满足官方 DTD 的结构要求（如权利要求书为空）。

    section 为出问题的 `## 章节` 标题，供上层定位行。
    """

    def __init__(self, section, message):
        super().__init__(message)
        self.section = section


@dataclass
class BookResult:
    code: str
    name: str
    file: str
    valid: bool
    errors: list = field(default_factory=list)
    images: list = field(default_factory=list)

    def to_dict(self):
        return {"code": self.code, "name": self.name, "file": self.file, "valid": self.valid,
                "errors": self.errors, "images": self.images}


@dataclass
class XmlExportResult:
    out_dir: str
    zip_path: str
    books: list
    warnings: list
    counts: dict

    @property
    def valid(self):
        return all(b.valid for b in self.books)

    def to_dict(self):
        return {"zip": os.path.basename(self.zip_path), "valid": self.valid,
                "books": [b.to_dict() for b in self.books], "warnings": self.warnings,
                "counts": self.counts}


# --------------------------------------------------------------------------- 行内文本
def inline_text(parent, text, sub_except=DEFAULT_SUB_EXCEPT, tail_target=None):
    """把 content.md 一段文本写入 parent：与 builder.add_math_text 同一 inline.inline_tokens，
    b/i/u/sub/sup → <b><i><u><sub><sup>（嵌套时按 MARKS 顺序由外向内；DTD 允许五者互相包含）。

    不写入 builder 为 Word 断行而插入的零宽字符。返回最后一个顶层行内元素（无则 tail_target）。"""
    text = _INLINE_MATH.sub(lambda m: latex_to_text(m.group(1) or m.group(2)), text)
    last = tail_target
    for marks, seg in inline_tokens(text, sub_except):
        if not marks:
            _append_text(parent, last, seg)
            continue
        outer = inner = None
        for m in MARKS:
            if m in marks:
                el = etree.SubElement(parent if inner is None else inner, m)
                outer = outer if outer is not None else el
                inner = el
        inner.text = seg
        last = outer
    return last


def _append_text(parent, last, seg):
    if not seg:
        return
    if last is None:
        parent.text = (parent.text or "") + seg
    else:
        last.tail = (last.tail or "") + seg


def claim_text_with_refs(ct, text, claim_num, sub_except):
    """claim-text：把"权利要求N"/"权利要求M至N"包成 <claim-ref idref="cM … cN">（仅引用更小权项号）。"""
    last = None
    pos = 0
    for m in _CLAIM_REF_RE.finditer(text):
        lo, hi = int(m.group(1)), int(m.group(2) or m.group(1))
        if not (1 <= lo <= hi < claim_num):
            continue
        last = inline_text(ct, text[pos:m.start()], sub_except, tail_target=last)
        ref = etree.SubElement(ct, "claim-ref", idref=" ".join(f"c{i}" for i in range(lo, hi + 1)))
        ref.text = m.group(0)
        last = ref
        pos = m.end()
    inline_text(ct, text[pos:], sub_except, tail_target=last)


# --------------------------------------------------------------------------- 图像资产
@dataclass
class Asset:
    kind: str  # fig | eq | tbl
    book: str
    name: str  # 输出 jpg 文件名（全案唯一）
    payload: object
    width_cm: float | None = None
    he_mm: int = 0
    wi_mm: int = 0
    el: object = None  # 对应的 <img>，栅格化后回填 he/wi


def _render_assets(assets, style, sub_except, figs_dir, docx_to_pdf, pdftoppm, workdir):
    """所有对象各占一页 → 一份 DOCX → PDF → 300 dpi JPG → 按内容裁边。就地填充 Asset.he_mm/wi_mm。"""
    if not assets:
        return
    if shutil.which(pdftoppm) is None:
        raise XmlExportError(f"缺少 {pdftoppm}（poppler-utils），无法把公式/表格/附图栅格化为 JPG")
    b = PatentBuilder(sub_except=sub_except, style=style)
    doc = b.doc
    b.a4(doc.sections[0])
    for i, a in enumerate(assets):
        if i:
            doc.add_page_break()
        if a.kind == "fig":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            path = os.path.join(figs_dir, a.payload)
            try:
                p.add_run().add_picture(path, width=Cm(a.width_cm))
            except Exception as e:  # noqa: BLE001
                raise FigureError(path, e) from e
        elif a.kind == "eq":
            b.equation(a.payload, "")
        else:
            rows, opts = a.payload
            b.add_table(rows, b.table_opts(opts))
    docx_path = os.path.join(workdir, "assets.docx")
    doc.save(docx_path)
    docx_to_pdf(docx_path, workdir)
    pdf_path = os.path.join(workdir, "assets.pdf")
    prefix = os.path.join(workdir, "page")
    subprocess.run([pdftoppm, "-r", str(DPI), "-jpeg", "-jpegopt", "quality=92", pdf_path, prefix],
                   check=True, capture_output=True, timeout=300)
    pages = sorted(f for f in os.listdir(workdir) if f.startswith("page") and f.endswith(".jpg"))
    crops = []
    for name in pages:
        img = Image.open(os.path.join(workdir, name)).convert("L")
        bbox = ImageChops.invert(img.point(lambda v: 255 if v > 200 else 0)).getbbox()
        if bbox is None:
            crops.append(None)
            continue
        pad = DPI // 25  # ~1 mm
        box = (max(bbox[0] - pad, 0), max(bbox[1] - pad, 0),
               min(bbox[2] + pad, img.width), min(bbox[3] + pad, img.height))
        crops.append(img.crop(box))
    if len(crops) == len(assets):
        blank = [i for i, c in enumerate(crops) if c is None]
        if blank:
            kinds = "、".join(sorted({ASSET_KIND_NAMES[assets[i].kind] for i in blank}))
            raise XmlExportError(f"{len(blank)} 个对象（{kinds}）渲染为空白页：LibreOffice 可能缺少 Math 组件"
                                 "（libreoffice-math）或对应字体，请检查转换环境")
    else:
        crops = [c for c in crops if c is not None]  # 表格/公式后 page break 产生的空白尾页
    if len(crops) != len(assets):
        raise XmlExportError(f"对象栅格化页数不符：期望 {len(assets)} 个对象，得到 {len(crops)} 页"
                             "（某个表格/公式/附图可能跨页，请缩小其宽度或拆分）")
    for a, im in zip(assets, crops):
        a.wi_mm = max(1, round(im.width / DPI * 25.4))
        a.he_mm = max(1, round(im.height / DPI * 25.4))
        im.convert("RGB").save(os.path.join(workdir, a.name), "JPEG", quality=92, dpi=(DPI, DPI))
        a.el.set("he", str(a.he_mm))
        a.el.set("wi", str(a.wi_mm))


def _img(parent, asset, content="drawing"):
    asset.el = etree.SubElement(parent, "img", he="0", wi="0", file=asset.name,
                                **{"img-content": content, "img-format": "jpg", "inline": "no"})
    return asset.el


# --------------------------------------------------------------------------- 五书装配
class _Exporter:
    def __init__(self, config):
        self.config = config
        self.sub_except = tuple(config.get("sub_except", DEFAULT_SUB_EXCEPT))
        self.style = merge_style(config.get("style"))
        self.assets = []
        self.warnings = []
        self.counters = {}
        self.dropped_inline_figs = 0

    def _asset(self, kind, book, payload, width_cm=None):
        tag = {"fig": "", "eq": "m", "tbl": "t"}[kind]
        key = (book, kind)
        self.counters[key] = self.counters.get(key, 0) + 1
        num = self.counters[key]
        a = Asset(kind, book, f"{book}-{tag}{num:04d}.jpg", payload, width_cm)
        self.assets.append(a)
        return a, num

    def _root(self, book_child):
        root = etree.Element("cn-application-body", lang="zh", country="CN")
        return root, etree.SubElement(root, book_child)

    def _maths(self, host, book, formula):
        a, num = self._asset("eq", book, formula)
        m = etree.SubElement(host, "maths", num=f"{num:04d}")
        _img(m, a, "character")

    # 100001
    def claims(self, claims):
        root, cc = self._root("cn-claims")
        for n, c in enumerate(claims, 1):
            cl = etree.SubElement(cc, "claim", id=f"c{n}", num=str(n))
            last_ct = None
            for seg in c.split("\n"):
                if seg.startswith("@EQ@"):
                    if last_ct is None:
                        last_ct = etree.SubElement(cl, "claim-text")
                    self._maths(last_ct, "100001", seg[4:])
                else:
                    last_ct = etree.SubElement(cl, "claim-text")
                    claim_text_with_refs(last_ct, seg, n, self.sub_except)
            if last_ct is None:
                etree.SubElement(cl, "claim-text")
        return root

    # 100002
    def description(self, title, spec):
        root, desc = self._root("description")
        etree.SubElement(desc, "invention-title").text = title
        host = desc  # 当前 ### 对应的 section 元素（或 desc）；#### 作为 level=2 heading 直挂其下
        n = 0
        last_p = None

        def num_p(text=None):
            nonlocal n, last_p
            n += 1
            p = etree.SubElement(host, "p", num=f"{n:04d}", Italic="0")
            if text is not None:
                inline_text(p, text, self.sub_except)
            last_p = p
            return p

        for item in spec:
            if isinstance(item, tuple):
                kind = item[0]
                if kind == "h":
                    el = SECTION_ELEMENTS.get(item[1])
                    host = etree.SubElement(desc, el) if el else desc
                    etree.SubElement(host, "heading", level="1").text = item[1]
                    last_p = None
                elif kind == "s":
                    etree.SubElement(host, "heading", level="2").text = item[1]
                    last_p = None
                elif kind == "e":
                    target = last_p if last_p is not None else num_p()
                    self._maths(target, "100002", item[1])
                elif kind == "t":
                    rows, opts = item[1], (item[2] if len(item) > 2 else {})
                    p = num_p()
                    a, tnum = self._asset("tbl", "100002", (rows, opts))
                    t = etree.SubElement(p, "tables", num=f"{tnum:04d}")
                    _img(t, a, "character")
                    if opts.get("表题后置"):
                        num_p(opts["表题后置"])
                    if opts.get("表注"):
                        num_p(opts["表注"])
                elif kind == "f":
                    # 说明书内嵌图（W18）：非递交版式元素，XML 包剔除，仅记跳过
                    self.dropped_inline_figs += 1
                    last_p = None
                elif kind == "c":
                    p = num_p()
                    pre = etree.SubElement(p, "pre")
                    pre.text = "\n".join(item[1])
            else:
                num_p(item)
        return root, n

    # 100003
    def drawings(self, spec_figures):
        root, dr = self._root("cn-drawings")
        for i, (fname, w) in enumerate(spec_figures, 1):
            a, _ = self._asset("fig", "100003", fname, w)
            # figure-labels 是官方预览在图下渲染的图号标签（showxml.xsl 取
            # ../@figure-labels），DOCX 附图书同位有「图N」题注；缺失则图下空槽
            fig = etree.SubElement(dr, "figure", id=f"f{i:04d}", num=str(i),
                                 **{"figure-labels": f"图{i}"})
            _img(fig, a)
        return root

    # 100004
    def abstract(self, text):
        root, ab = self._root("cn-abstract")
        for i, para in enumerate([s for s in text.split("\n") if s.strip()], 1):
            p = etree.SubElement(ab, "p", num=f"{i:04d}", Italic="0")
            inline_text(p, para, self.sub_except)
        return root

    # 100005
    def abstract_figure(self, fname, width):
        root, ab = self._root("cn-abstract")
        af = etree.SubElement(ab, "cn-abst-figure")
        a, _ = self._asset("fig", "100005", fname, width)
        fig = etree.SubElement(af, "figure", id="fa0001", num="1")
        _img(fig, a)
        return root


def serialize(root):
    body = etree.tostring(root, encoding="unicode", pretty_print=True)
    return f'<?xml version="1.0" encoding="UTF-8"?>\n{STYLESHEET_PI}\n{DOCTYPE}\n{body}'


_dtd_cache = {}


def load_dtd(path=DTD_PATH):
    if path not in _dtd_cache:
        _dtd_cache[path] = etree.DTD(path)
    return _dtd_cache[path]


def validate_xml(path):
    """官方 DTD 校验；返回 (valid, errors)。DTD 自身非确定性内容模型的报错不计入。"""
    parser = etree.XMLParser(load_dtd=False, resolve_entities=False, no_network=True)
    try:
        tree = etree.parse(path, parser)
    except etree.XMLSyntaxError as e:
        return False, [str(e)]
    dtd = load_dtd()
    dtd.validate(tree)
    errors = [str(e.message) for e in dtd.error_log if NONDETERMINISTIC_DTD_MSG not in e.message]
    return not errors, errors


def gb18030_issues(root):
    bad = set()
    for t in root.itertext():
        for ch in t:
            try:
                ch.encode("gb18030")
            except UnicodeEncodeError:
                bad.add(ch)
    return sorted(bad)


def export_cnipa_xml(config, out_dir, docx_to_pdf, pdftoppm="pdftoppm", zip_name="cnipa-xml.zip"):
    """按 build_patent 的 config 导出五书 XML + JPG 到 out_dir/<code>/，并打成 out_dir/zip_name。

    返回 XmlExportResult；结构性问题（DTD 不通过、GB18030 越界字符）记入 books[].errors / warnings，
    渲染失败（附图缺失、对象跨页、pdftoppm 缺失）抛异常。"""
    ex = _Exporter(config)
    title, abstract, claims, spec = load_md(config["content_md"])
    if not claims:
        raise XmlContentError("权利要求书", "权利要求书为空：官方 DTD 要求至少 1 项权利要求，请先补写权利要求")
    roots = {}
    roots["100001"] = ex.claims(claims)
    roots["100002"], paragraphs = ex.description(title, spec)
    spec_figures = config.get("spec_figures") or []
    if spec_figures:
        roots["100003"] = ex.drawings(spec_figures)
    roots["100004"] = ex.abstract(abstract)
    if config.get("abstract_figure"):
        af, aw = config["abstract_figure"]
        roots["100005"] = ex.abstract_figure(af, aw)
    if not spec_figures:
        ex.warnings.append("未配置说明书附图，未生成 100003（说明书附图）")
    if ex.dropped_inline_figs:
        ex.warnings.append(f"说明书内含 {ex.dropped_inline_figs} 处内嵌图（非递交版式），已从 XML 剔除")
    codes = [(c, n) for c, n in BOOKS if c in roots]

    os.makedirs(out_dir, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="cnipa-xml-") as tmp:
        _render_assets(ex.assets, ex.style, ex.sub_except, config["figs_dir"], docx_to_pdf,
                       pdftoppm, tmp)
        for code, _ in codes:
            d = os.path.join(out_dir, code)
            if os.path.isdir(d):
                shutil.rmtree(d)
            os.makedirs(d)
            for a in ex.assets:
                if a.book == code:
                    shutil.copy(os.path.join(tmp, a.name), os.path.join(d, a.name))

    books = []
    for code, name in codes:
        path = os.path.join(out_dir, code, f"{code}.xml")
        with open(path, "w", encoding="utf-8") as f:
            f.write(serialize(roots[code]))
        valid, errors = validate_xml(path)
        bad = gb18030_issues(roots[code])
        if bad:
            ex.warnings.append(f"{code} 含 GB18030 之外字符：{' '.join(bad)}（官方工具会要求转图）")
        books.append(BookResult(code, name, f"{code}/{code}.xml", valid, errors,
                                [a.name for a in ex.assets if a.book == code]))

    zip_path = os.path.join(out_dir, zip_name)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for code, _ in codes:
            d = os.path.join(out_dir, code)
            for fn in sorted(os.listdir(d)):
                z.write(os.path.join(d, fn), f"{code}/{fn}")
    counts = {"claims": len(claims), "paragraphs": paragraphs,
              "figures": len(spec_figures),
              "equations": sum(a.kind == "eq" for a in ex.assets),
              "tables": sum(a.kind == "tbl" for a in ex.assets),
              "abstract_chars": len(abstract)}
    return XmlExportResult(out_dir, zip_path, books, ex.warnings, counts)
