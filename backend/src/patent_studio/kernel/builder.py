# -*- coding: utf-8 -*-
"""CN 发明专利申请文件通用排版装配器。

组织：说明书摘要 / 摘要附图 / 权利要求书（独立编页）/ 说明书（独立编页）/ 说明书附图
格式：A4；正文宋体小四两端对齐；标题黑体；说明书段落 [0001] 自动编号；
     行内下标记号（_x / _{...} / _[...]）渲染为 Word 真下标；
     $$...$$ 独立公式渲染为 OMML 原生 Word 公式对象（Word 中可双击编辑）。
依赖：python-docx、lxml。

全部排版常量集中在 style.DEFAULT_STYLE，可通过 config["style"] 部分覆盖；
不覆盖时输出与原硬编码实现逐字节一致。
"""
import os
import re
import struct
from dataclasses import dataclass

import docx.image as _dimg
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import BaseImageHeader
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from .content import load_md
from .inline import SUB_RE, inline_tokens
from .numbering import number_spec, para_label
from .omml import append_omml, latex_flat_inline
from .style import DEFAULT_STYLE, heading_prefix, heading_style, merge_style

_INLINE_MATH = re.compile(r"\$\$([^$\n]+?)\$\$|\$([^$\n]+?)\$")

# style.headings.*.align → Word 对齐；left 即 builder 原来的两端对齐（单行标题视觉相同，零漂移逐字节不变）
_HEADING_ALIGN = {"left": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER}


# --- EMF 矢量图支持（python-docx 原生不识别 EMF 格式）---
class _Emf(BaseImageHeader):
    @classmethod
    def from_stream(cls, stream):
        stream.seek(24)
        l, t, r, b = struct.unpack("<4i", stream.read(16))
        dpi = 96
        return cls(int((r - l) / 100.0 / 25.4 * dpi),
                   int((b - t) / 100.0 / 25.4 * dpi), dpi, dpi)

    @property
    def content_type(self):
        return "image/x-emf"

    @property
    def default_ext(self):
        return "emf"


_dimg.SIGNATURES = tuple(_dimg.SIGNATURES) + ((_Emf, 0, b"\x01\x00\x00\x00"),)


def probe_image(data):
    """用 add_picture 同一套识图器（含上面注册的 EMF）读图头；不是图就抛识图器的原异常。"""
    return _dimg.image.Image.from_blob(data)


SONG, HEI = DEFAULT_STYLE["fonts"]["body_cjk"], DEFAULT_STYLE["fonts"]["heading_cjk"]

DEFAULT_SUB_EXCEPT = ("tsf_guard", "on_matured")
_SUB_RE = SUB_RE  # 下标记号（W9 起与 inline.py 同源；保留名字供 cnipa_xml / docx_import / omml 引用）


class EquationError(Exception):
    """独立公式解析/渲染失败；formula 为原式文本，供上层定位 content.md 行号。"""

    def __init__(self, formula, cause):
        super().__init__("公式解析失败：%s（%s: %s）" % (formula, type(cause).__name__, cause))
        self.formula = formula
        self.cause = cause


class FigureError(Exception):
    """附图文件缺失或无法读取。"""

    def __init__(self, path, cause):
        super().__init__("附图无法读取：%s（%s）" % (os.path.basename(path), cause))
        self.path = path
        self.cause = cause


@dataclass
class BuildResult:
    out: str
    abstract_chars: int
    claims: int
    paragraphs: int
    title: str


def set_font(run, name=SONG, size=12, bold=False, ascii_name=None):
    run.font.name = ascii_name or "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element.rPr
    rf = r.rFonts if r.rFonts is not None else r.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), name)


class PatentBuilder:
    def __init__(self, sub_except=DEFAULT_SUB_EXCEPT, style=None, run_hook=None):
        self.sub_except = tuple(sub_except)
        self.style = style if style is not None else merge_style(None)
        self.run_hook = run_hook
        self.doc = Document()

    # --- style 快捷访问 ---
    @property
    def song(self):
        return self.style["fonts"]["body_cjk"]

    @property
    def hei(self):
        return self.style["fonts"]["heading_cjk"]

    def sz(self, key):
        return self.style["sizes"][key]

    def gap(self, key):
        return self.style["paragraph"]["space_after"][key]

    def _font(self, run, name, size, bold=False, label=False):
        """set_font 的 style 感知版：西文字体取自 style。"""
        fonts = self.style["fonts"]
        set_font(run, name, size, bold,
                 ascii_name=fonts["label_ascii"] if label else fonts["ascii"])

    # --- 文本与公式 ---
    def add_math_text(self, p, text, font=None, size=None, bold=False):
        """按下标记号切分文本：_x / _{...} / _[...] 渲染为 Word 真下标。

        run_hook（patent_studio.kernel.tracked.RunHook）不为 None 时，先由 hook.split 把文本切成
        (kind, 片段) 序列，每片段照常生成 run 后交给 hook.wrap 包裹（w:ins / w:del 等）。
        """
        if self.run_hook is None:
            self._emit_math_text(p, text, font, size, bold)
            return
        for kind, seg in self.run_hook.split(text):
            n = len(p._p)
            self._emit_math_text(p, seg, font, size, bold)
            self.run_hook.wrap(p, kind, list(p._p)[n:])

    def _emit_math_text(self, p, text, font, size, bold):
        font = font or self.song
        size = self.sz("body") if size is None else size
        pos = 0
        for mm in _INLINE_MATH.finditer(text):
            if mm.start() > pos:
                self._emit_plain_text(p, text[pos:mm.start()], font, size, bold)
            eq = mm.group(1) or mm.group(2)
            try:
                flat = latex_flat_inline(eq)
            except Exception:
                flat = None
            if flat is not None:
                # 简单行内式（字母/上下标/符号）：vertAlign 文本排版，不撑行距
                self._emit_plain_text(p, flat, font, size, bold)
                pos = mm.end()
                continue
            try:
                append_omml(p, eq, size, inline=True)
            except Exception:  # 公式解析失败不破坏段落：原样按文本落
                self._emit_plain_text(p, mm.group(0), font, size, bold)
            pos = mm.end()
        if pos < len(text):
            self._emit_plain_text(p, text[pos:], font, size, bold)

    def _emit_plain_text(self, p, text, font, size, bold):
        text = re.sub(r"(?<=[0-9]),(?=[0-9.])", ",\u200b", text)
        text = re.sub(r"(?<=[、，；：])(?=[A-Za-z0-9])", "\u200b", text)
        text = re.sub(r"(?<=[A-Za-z0-9₀-₉η√])/(?=[A-Za-z0-9√η(])", "/\u2060", text)
        for marks, seg in inline_tokens(text, self.sub_except):
            r = p.add_run(seg)
            self._font(r, font, size, bold or "b" in marks)
            if "i" in marks:
                r.font.italic = True
            if "u" in marks:
                r.font.underline = True
            if "sub" in marks:
                r.font.subscript = True
            if "sup" in marks:
                r.font.superscript = True

    def para(self, text, size=None, bold=False, font=None,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, first_indent=None,
             keep_with_next=False):
        p = self.doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = self.style["paragraph"]["line_spacing"]
        if keep_with_next:
            p.paragraph_format.keep_with_next = True
        if first_indent:
            p.paragraph_format.first_line_indent = Pt(first_indent)
        self.add_math_text(p, text, font or self.song,
                           self.sz("body") if size is None else size, bold)
        return p

    def equation(self, text, num):
        """独立成行居中公式，右侧编号与公式整体垂直居中；
        公式为 OMML 原生 Word 公式对象。以无框三列表实现：
        左空列 = 右编号列，使公式位于版心正中。"""
        width = (self.doc.sections[-1].page_width
                 - self.doc.sections[-1].left_margin
                 - self.doc.sections[-1].right_margin)
        side = Cm(self.style["equation"]["side_cm"])
        lside, rside = side, side
        tbl = self.doc.add_table(rows=1, cols=3)
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right",
                     "insideH", "insideV"):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:val"), "none")
            borders.append(e)
        tblPr.append(borders)
        widths = (lside, width - lside - rside, rside)
        cell_gap = Pt(self.gap("equation_cell"))
        for col, w in zip(tbl.columns, widths):
            col.width = w
        for cell, w in zip(tbl.rows[0].cells, widths):
            cell.width = w
            va = OxmlElement("w:vAlign")
            va.set(qn("w:val"), "center")
            cell._tc.get_or_add_tcPr().append(va)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = cell_gap
            cp.paragraph_format.space_after = cell_gap
            cp.paragraph_format.line_spacing = self.style["paragraph"]["line_spacing"]
        mid = tbl.rows[0].cells[1].paragraphs[0]
        mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            append_omml(mid, text, self.sz("equation"))
        except Exception as e:  # noqa: BLE001 - 任何解析失败都带上原式抛出
            raise EquationError(text, e) from e
        right = tbl.rows[0].cells[2].paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = right.add_run(re.sub(r"^式\s*", "", num))
        self._font(run, self.song, self.sz("equation"))
        return tbl

    def heading(self, text, size=15):
        return self.para(text, size=size, bold=True, font=self.hei,
                         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    PART_CODES = {"权利要求书": "100001", "说明书": "100002",
                  "说明书附图": "100003", "说明书摘要": "100004",
                  "摘要附图": "100005"}

    def part_heading(self, section, part):
        """部分标题：写入节页眉（每页重复），黑体四号、字间加宽，
        页脚写表格编号+年份页尾
        （对齐官方 2023.03 模板 100001–100005）。

        part 为 style.PART_KEYS 之一（abstract/abstract_figure/claims/description/figures）；
        为兼容旧调用也接受中文标签本身。"""
        hs = self.style["header"]
        labels = hs["labels"]
        if part not in labels:
            part = next((k for k, v in labels.items() if v == part), part)
        text = labels[part]
        section.header.is_linked_to_previous = False
        hp = section.header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = hp._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(hs["rule_sz"]))
        bottom.set(qn("w:space"), str(hs["rule_space"]))
        bottom.set(qn("w:color"), "000000")
        pbdr.append(bottom)
        pPr.append(pbdr)
        run = hp.add_run(text)
        self._font(run, self.hei, self.sz("header"), label=True)
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(hs["spacing_short"] if len(text) <= 3
                                else hs["spacing_long"]))
        run._r.rPr.append(sp)
        self.part_footer(section, self.style["footer"]["codes"][part])
        return hp

    def part_footer(self, section, code):
        """页尾：上横线，左侧表格编号与年份，页码居中
        （对齐官方 2023.03 模板页尾）。"""
        fs = self.style["footer"]
        section.footer_distance = Cm(self.style["page"]["footer_distance_cm"])
        pgp = section.footer.paragraphs[0]
        p1 = pgp.insert_paragraph_before()
        pPr = p1._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), str(fs["rule_sz"]))
        top.set(qn("w:space"), str(fs["rule_space"]))
        top.set(qn("w:color"), "000000")
        pbdr.append(top)
        pPr.append(pbdr)
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(code)
        self._font(r1, self.hei, self.sz("footer"), label=True)
        p2 = pgp.insert_paragraph_before()
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(fs["year_text"])
        self._font(r2, self.hei, self.sz("footer"), label=True)

    def num_para(self, num, text="", size=None, space_after=6, keep_with_next=False):
        """说明书编号段：段号[000n]黑体小五、紧接正文（对齐公告版段号版式）。"""
        size = self.sz("body") if size is None else size
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = self.style["paragraph"]["line_spacing"]
        if keep_with_next:
            p.paragraph_format.keep_with_next = True
        run = p.add_run(num)
        self._font(run, self.hei, self.sz("para_number"), label=True)
        if text:
            gap = p.add_run("\u3000")
            self._font(gap, self.song, size)
            self.add_math_text(p, text, self.song, size)
        return p

    # --- 版面 ---
    def a4(self, section):
        pg = self.style["page"]
        section.page_width = Cm(pg["width_cm"])
        section.page_height = Cm(pg["height_cm"])
        section.top_margin = Mm(pg["margin_top_mm"])
        section.bottom_margin = Mm(pg["margin_bottom_mm"])
        section.left_margin = Mm(pg["margin_left_mm"])
        section.right_margin = Mm(pg["margin_right_mm"])
        section.header_distance = Cm(pg["header_distance_cm"])
        sectPr = section._sectPr
        grid = sectPr.find(qn("w:docGrid"))
        if grid is None:
            grid = OxmlElement("w:docGrid")
            sectPr.append(grid)
        grid.set(qn("w:type"), "lines")
        grid.set(qn("w:linePitch"), str(pg["line_pitch"]))
        grid.set(qn("w:charSpace"), "0")

    @staticmethod
    def restart_page_number(section):
        sectPr = section._sectPr
        pg = sectPr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sectPr.append(pg)
        pg.set(qn("w:start"), "1")

    def add_page_field(self, footer_para):
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_para.add_run()
        for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                                 ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
                                 ("w:fldChar", {"w:fldCharType": "end"}, None)):
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            if text:
                el.text = text
            run._r.append(el)
        self._font(run, self.song, self.sz("page_number"))

    def new_part(self, restart=True):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self.a4(section)
        section.footer.is_linked_to_previous = False
        if restart:
            self.restart_page_number(section)
        fp = section.footer.paragraphs[0]
        fp.text = ""
        self.add_page_field(fp)
        return section

    @staticmethod
    def _set_borders(el_pr, tag, spec):
        """在 tblPr/tcPr 下写入边框定义。spec: {方位: sz 或 None(无线)}。"""
        borders = el_pr.find(qn("w:" + tag))
        if borders is None:
            borders = OxmlElement("w:" + tag)
            el_pr.append(borders)
        for side, sz in spec.items():
            b = OxmlElement("w:" + side)
            if sz is None:
                b.set(qn("w:val"), "none")
            else:
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), str(sz))
                b.set(qn("w:color"), "000000")
            borders.append(b)

    def table_opts(self, opts):
        """content.md 声明的表格选项叠加在 style.table.defaults 之上。"""
        merged = {k: v for k, v in self.style["table"]["defaults"].items() if v}
        merged.update(opts or {})
        return merged

    def add_table(self, rows, opts=None, container=None):
        """说明书表格；默认全框细黑线（对齐授权案主流版式），首行为表头。

        opts（均可选）：三线 / 重复表头 / 列宽=[百分比…]；单元格 "^^" 与上行竖向合并。"""
        opts = opts or {}
        ts = self.style["table"]
        host = container if container is not None else self.doc
        ncols = max(len(r) for r in rows)
        tbl = host.add_table(rows=len(rows), cols=ncols)
        if opts.get("三线"):
            tbl.style = None
            tblPr = tbl._tbl.tblPr
            self._set_borders(tblPr, "tblBorders",
                              {"top": ts["rule_outer_sz"], "bottom": ts["rule_outer_sz"],
                               "left": None, "right": None, "insideH": None,
                               "insideV": None})
        else:
            tbl.style = "Table Grid"
        tbl.alignment = 1  # 居中
        if opts.get("列宽"):
            tblPr = tbl._tbl.tblPr
            tw = OxmlElement("w:tblW")
            tw.set(qn("w:type"), "pct")
            tw.set(qn("w:w"), "5000")
            tblPr.append(tw)
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tblPr.append(layout)
        for i, r in enumerate(rows):
            if i == 0 and opts.get("重复表头"):
                trPr = tbl.rows[0]._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            for j in range(ncols):
                cell = tbl.cell(i, j)
                txt = r[j] if j < len(r) else ""
                tcPr = cell._tc.get_or_add_tcPr()
                if opts.get("列宽") and j < len(opts["列宽"]):
                    cw = OxmlElement("w:tcW")
                    cw.set(qn("w:type"), "pct")
                    cw.set(qn("w:w"), str(int(opts["列宽"][j] * 50)))
                    tcPr.append(cw)
                if txt == "^^":
                    vm = OxmlElement("w:vMerge")
                    tcPr.append(vm)  # 无 val = continue
                    txt = ""
                elif i + 1 < len(rows) and j < len(rows[i + 1]) \
                        and rows[i + 1][j] == "^^":
                    vm = OxmlElement("w:vMerge")
                    vm.set(qn("w:val"), "restart")
                    tcPr.append(vm)
                vAlign = OxmlElement("w:vAlign")
                vAlign.set(qn("w:val"), "center")
                tcPr.append(vAlign)
                if opts.get("三线") and i == 0:
                    self._set_borders(tcPr, "tcBorders", {"bottom": ts["rule_head_sz"]})
                cell.paragraphs[0].text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(ts["cell_space_after"])
                if i == 0:
                    p.paragraph_format.keep_with_next = True
                self.add_math_text(p, txt, self.song, self.sz("table"), bold=(i == 0))
        return tbl

    def add_table_with_left_num(self, num_text, rows, opts=None):
        """段号居左版式：外层无框线 1×2 表，左格段号纵向居中，右格内嵌数据表。"""
        outer = self.doc.add_table(rows=1, cols=2)
        outer.style = None
        outer.alignment = 1
        tblPr = outer._tbl.tblPr
        self._set_borders(tblPr, "tblBorders",
                          {"top": None, "bottom": None, "left": None,
                           "right": None, "insideH": None, "insideV": None})
        tw = OxmlElement("w:tblW")
        tw.set(qn("w:type"), "dxa")
        content_width = (self.doc.sections[-1].page_width
                         - self.doc.sections[-1].left_margin
                         - self.doc.sections[-1].right_margin)
        shell = min(Cm(15.4), content_width)  # 上限 15.4cm；自定义页边距收窄版心时随版心收缩
        tw.set(qn("w:w"), str(int(shell.twips)))
        tblPr.append(tw)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        trPr = outer.rows[0]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        lcell, rcell = outer.cell(0, 0), outer.cell(0, 1)
        lcell.width = Cm(1.9)
        rcell.width = shell - lcell.width
        tcPr = lcell._tc.get_or_add_tcPr()
        vAlign = OxmlElement("w:vAlign")
        vAlign.set(qn("w:val"), "center")
        tcPr.append(vAlign)
        p = lcell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(num_text)
        self._font(run, self.hei, self.sz("para_number"), label=True)
        rcell.paragraphs[0].text = ""
        opts = dict(opts or {})
        if not opts.get("列宽"):
            ncols = max(len(r) for r in rows)
            opts["列宽"] = [100.0 / ncols] * ncols
        self.add_table(rows, opts, container=rcell)
        return outer

    def add_pseudocode(self, lines_, num_text):
        """伪代码段：一个编号段内多行等宽体（对齐授权案伪代码入编号段写法）。"""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(self.gap("body"))
        p.paragraph_format.line_spacing = self.style["paragraph"]["line_spacing"]
        run = p.add_run(num_text)
        self._font(run, self.hei, self.sz("para_number"), label=True)
        for k, ln in enumerate(lines_):
            if k:
                p.add_run().add_break()
            run = p.add_run(ln)
            run.font.name = self.style["fonts"]["mono"]
            run.font.size = Pt(self.sz("code"))
            run._element.rPr.rFonts.set(qn("w:eastAsia"), self.song)
        return p

    def add_figure(self, path, caption, width_cm=15.0, max_height_cm=19.5):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        try:
            with open(path, "rb") as fh:
                data = fh.read()
            # 图+题注须同页（keepNext 已写）——图块高到容不下题注时该约束不可满足，
            # LO/Word 退化为题注孤行翻到次页（bert-12 图4 实案）。限高取纵高优先，
            # 保证「图+题注+间距」≤ 一页内容区（A4 约 23.7cm 可用）。
            img = probe_image(data)
            tall = img.px_height * width_cm > img.px_width * max_height_cm
            p.add_run().add_picture(path, height=Cm(max_height_cm) if tall else None,
                                  width=None if tall else Cm(width_cm))
        except Exception as e:  # noqa: BLE001 - 缺文件/格式不识别都归为附图错误
            raise FigureError(path, e) from e
        self.para(caption, align=WD_ALIGN_PARAGRAPH.CENTER,
                  space_after=self.gap("figure_caption"))


def build_patent_ex(config):
    """build_patent 的详细版：返回 BuildResult（含摘要字数/权项数/段落数），不打印。

    config["abstract_figure"] = (file, width_cm)，config["spec_figures"] = [(file, width_cm), …]：
    Studio 里由 ProjectStore.build_config 从 docs/figures.md（figures_md.FiguresDoc.legacy_view）派生，
    不再来自 studio.json；旧脚本（examples/*/build_docx.py CONFIG）仍直接给元组，zero_drift 三路必须同 md5。
    图号按顺序生成（图1..N），alt 不进 DOCX。"""
    style = merge_style(config.get("style"))
    b = PatentBuilder(sub_except=config.get("sub_except", DEFAULT_SUB_EXCEPT), style=style,
                      run_hook=config.get("run_hook"))
    title, abstract, claims, spec = load_md(config["content_md"])
    figs = config["figs_dir"]
    doc = b.doc
    gap = b.gap

    st = doc.settings.element
    tab = st.find(qn("w:defaultTabStop"))
    if tab is None:
        tab = OxmlElement("w:defaultTabStop")
        st.append(tab)
    tab.set(qn("w:val"), str(style["page"]["default_tab_stop"]))

    b.a4(doc.sections[0])
    doc.sections[0].footer.is_linked_to_previous = False
    fp = doc.sections[0].footer.paragraphs[0]
    b.add_page_field(fp)

    first_indent = style["paragraph"]["first_indent_pt"]
    b.part_heading(doc.sections[0], "abstract")
    b.para(title, size=b.sz("title"), font=b.hei, align=WD_ALIGN_PARAGRAPH.CENTER,
           space_after=gap("title"))
    limit = config.get("abstract_limit", 300)
    assert len(abstract) <= limit, f"摘要超过{limit}字：{len(abstract)}"
    b.para(abstract, first_indent=first_indent, space_after=gap("body"))

    if config.get("abstract_figure"):  # 纯方法无附图案卷省略摘要附图部
        b.part_heading(b.new_part(), "abstract_figure")
        af, aw = config["abstract_figure"]
        b.add_figure(os.path.join(figs, af), "图1", width_cm=aw)

    b.part_heading(b.new_part(restart=True), "claims")
    for c in claims:
        segs = c.split("\n")
        for i, seg in enumerate(segs):
            if seg.startswith("@EQ@"):
                b.equation(seg[4:], "")
            else:
                b.para(seg, first_indent=first_indent,
                       space_after=gap("claim_inner") if i < len(segs) - 1
                       else gap("claim_last"))

    b.part_heading(b.new_part(restart=True), "description")
    b.para(title, size=b.sz("title"), font=b.hei, align=WD_ALIGN_PARAGRAPH.CENTER,
           space_after=0)
    b.para("", space_after=0)
    numbers = number_spec(spec)
    sec_no, sub_no = 0, 0
    for idx, item in enumerate(spec):
        nums = numbers[idx]
        next_is_table = (idx + 1 < len(spec)
                         and isinstance(spec[idx + 1], tuple)
                         and spec[idx + 1][0] == "t")
        if isinstance(item, tuple):
            if item[0] == "f":
                b.add_figure(os.path.join(figs, item[1]), item[3] or "", item[2])
                continue
            if item[0] == "e":
                b.equation(item[1], item[2])
                continue
            if item[0] == "t":
                topts = b.table_opts(item[2] if len(item) > 2 else {})
                if topts.get("段号居左"):
                    b.add_table_with_left_num(para_label(nums.para_no), item[1], topts)
                else:
                    b.num_para(para_label(nums.para_no), space_after=gap("table_label"),
                               keep_with_next=True)
                    b.add_table(item[1], topts)
                if topts.get("表题后置"):
                    b.num_para(para_label(nums.caption_after_no), topts["表题后置"],
                               space_after=gap("table_label"))
                if topts.get("表注"):
                    b.num_para(para_label(nums.note_no), topts["表注"], size=b.sz("table"),
                               space_after=gap("table_label"))
                b.para("", space_after=gap("table_gap"))
                continue
            if item[0] == "c":
                b.add_pseudocode(item[1], para_label(nums.para_no))
                continue
            kind, text = item
            level = "section" if kind == "h" else "subsection"
            if kind == "h":
                sec_no, sub_no = sec_no + 1, 0
            else:
                sub_no += 1
            hs = heading_style(style, level)
            prefix = heading_prefix(hs["numbering"], sec_no, None if kind == "h" else sub_no)
            b.para(prefix + text, size=hs["size"], font=hs["font"], bold=hs["bold"],
                   align=_HEADING_ALIGN[hs["align"]], space_after=gap(level + "_heading"),
                   keep_with_next=next_is_table)
        else:
            b.num_para(para_label(nums.para_no), item, space_after=gap("body"),
                       keep_with_next=next_is_table)

    if config.get("spec_figures"):  # 无附图案卷省略说明书附图部
        b.part_heading(b.new_part(), "figures")
        for i, (fname, w) in enumerate(config["spec_figures"], 1):
            b.add_figure(os.path.join(figs, fname), f"图{i}", w)

    doc.save(config["out"])
    n = max((x for nm in numbers for x in (nm.para_no, nm.caption_after_no, nm.note_no) if x),
            default=0)
    return BuildResult(out=config["out"], abstract_chars=len(abstract),
                       claims=len(claims), paragraphs=n, title=title)


def build_patent(config):
    """按 config 生成一份完整的发明专利申请 docx。

    config 字段：
      content_md      content.md 路径（内容源，格式约定见 patent_studio.kernel.content）
      out             输出 docx 路径
      figs_dir        附图目录
      abstract_figure (文件名, 宽度cm)                摘要附图
      spec_figures    [(文件名, 宽度cm), ...]         说明书附图（按图1..N 顺序）
      sub_except      可选，不作下标解析的字面标识符元组
      abstract_limit  可选，摘要字数上限（默认 300）
      style           可选，排版参数覆盖（见 style.DEFAULT_STYLE；不传=原版式）
    """
    r = build_patent_ex(config)
    print("saved", r.out, " 摘要字数:", r.abstract_chars,
          " 权利要求项数:", r.claims, " 说明书段落数:", r.paragraphs)
    return r.out
