# -*- coding: utf-8 -*-
"""意见陈述书 DOCX（docs/plan/09 §3.4，python-docx 简版）。

export_docx(project, items, answers, claims_md) -> bytes
  标题 / 申请号占位 / 答复表（序号 · 审查意见 · 申请人意见 · 修改说明）/ 逐条陈述 / 修改后的权利要求书
  （当前权利要求书 + 已生成的 CriticMarkup 建议全部 accept 后的文本）。
"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from patent_docx import critic

FONT = "宋体"


def _font(run, size: float, bold: bool = False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _para(doc, text: str, size: float = 12, bold: bool = False, align=None, style: str | None = None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    _font(p.add_run(text), size, bold)
    return p


def _cell(cell, text: str, size: float = 10.5, bold: bool = False):
    cell.text = ""
    lines = text.split("\n") if text else [""]
    _font(cell.paragraphs[0].add_run(lines[0]), size, bold)
    for ln in lines[1:]:
        _font(cell.add_paragraph().add_run(ln), size, bold)


def _borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def amended_claims(claims_md: str, answers: list[dict]) -> str:
    """当前权利要求书源码 + 所有建议 accept 后的文本（按块替换，建议未命中原块则跳过）。"""
    text = claims_md
    for ans in answers:
        for e in ans.get("claim_edits") or []:
            if e.get("original") and e["original"] in text:
                text = text.replace(e["original"], critic.accept(e["critic_line"]), 1)
    return critic.accept(text).strip()


def export_docx(*, project: dict, items: list[dict], answers: list[dict], claims_md: str = "") -> bytes:
    by_id = {a["item_id"]: a for a in answers}
    doc = Document()
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    _para(doc, "意见陈述书", 22, True, WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"申请号：{project.get('app_no') or '________________'}", 12)
    _para(doc, f"发明名称：{project.get('title') or ''}", 12)
    _para(doc, "申请人针对审查意见通知书陈述意见如下：", 12)

    _para(doc, "一、答复表", 14, True)
    table = doc.add_table(rows=1, cols=4)
    _borders(table)
    for cell, head in zip(table.rows[0].cells, ("序号", "审查意见", "申请人意见", "修改说明")):
        _cell(cell, head, 10.5, True)
    for k, item in enumerate(items, 1):
        ans = by_id.get(item["id"]) or {}
        row = table.add_row().cells
        _cell(row[0], str(k))
        _cell(row[1], item["text"])
        _cell(row[2], "\n".join(f"{i}. {p}" for i, p in enumerate(ans.get("points") or [], 1)) or "（尚未答复）")
        _cell(row[3], ans.get("amend_summary") or ("无修改" if ans else ""))

    _para(doc, "二、逐条陈述", 14, True)
    for k, item in enumerate(items, 1):
        ans = by_id.get(item["id"]) or {}
        nums = "、".join(map(str, item.get("claims") or []))
        head = f"{k}. 关于" + (f"权利要求{nums}" if nums else "该项意见")
        _para(doc, head, 12, True)
        for p in ans.get("points") or []:
            _para(doc, p, 12)
        if ans.get("amend_summary"):
            _para(doc, "修改说明：" + ans["amend_summary"], 12)
        for e in ans.get("claim_edits") or []:
            _para(doc, f"权利要求 {e['claim']} 修改为：", 12)
            _para(doc, critic.accept(e["critic_line"]), 12)

    if claims_md.strip():
        _para(doc, "三、修改后的权利要求书", 14, True)
        for block in amended_claims(claims_md, answers).split("\n\n"):
            if block.strip():
                _para(doc, block.strip(), 12)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
