# -*- coding: utf-8 -*-
r"""LaTeX→OMML 排版回归（轮36）：s1 产出公式普遍是 LaTeX 记号，
此前线性解析器不识 \frac/\sum/\mathbb/\tag 等，源码原样印进 DOCX。"""

import io

import pytest
from docx import Document
from lxml import etree

from patent_studio.kernel.omml import append_omml, latex_to_text, to_mathml

M = {"m": "http://www.w3.org/1998/Math/MathML"}


def _mml(src):
    return etree.fromstring(to_mathml(src))


def _cnt(root, tag):
    return int(root.xpath(f"count(//m:{tag})", namespaces=M))


def test_frac_sup_nary_mathbb_tag():
    r = _mml(r"I\in \mathbb{R}^{H\times W\times C}\tag{6}")
    assert _cnt(r, "msup") == 1
    assert "ℝ" in latex_to_text(r"I\in \mathbb{R}^{H\times W\times C}\tag{6}")
    assert "tag" not in latex_to_text(r"a\tag{1}b")

    r = _mml(r"\sum_{i=1}^{N} x_i")
    assert _cnt(r, "munderover") == 1

    r = _mml(r"\frac{a+b}{c+d}")
    assert _cnt(r, "mfrac") == 1


def test_scripts_and_group_base():
    assert _cnt(_mml("b_i^2"), "msubsup") == 1
    assert _cnt(_mml("x^2_i"), "msubsup") == 1
    # 定界符组基（S_l(X^{(l)})_{p,q,c}）：sub 挂组、组内再含 sup
    r = _mml("S_l(X^{(l)})_{p,q,c}")
    assert _cnt(r, "msub") == 2 and _cnt(r, "msup") == 1
    assert _cnt(_mml(r"e^{x_i}"), "msup") == 1
    # 希腊字母等非字母数字下标
    assert _cnt(_mml("p_θ"), "msub") == 1


def test_latex_dropped_cmds_no_residue():
    out = latex_to_text(r"\mathcal{N}(\mu_\theta,\Sigma_\theta)")
    assert "\\" not in out and "mathcal" not in out
    assert "𝒩" in out


@pytest.mark.parametrize(
    "src",
    [
        r"\frac{a}{b}",
        r"x_i^2",
        r"\begin{pmatrix}a&b\\c&d\end{pmatrix}",
        r"\sqrt[3]{x}+1",
        r"\operatorname{softmax}(x)",
        r"a\cdot b\circ c",
    ],
)
def test_no_crash_no_backslash(src):
    out = latex_to_text(src)
    assert "\\" not in out
    assert out.strip()


def test_norm_bars_not_literal():
    out = latex_to_text(r"\lVert x_i - x_j \rVert_2^2")
    assert "Vert" not in out and "vert" not in out and "‖" in out
    assert "∣" in latex_to_text(r"\vert x \vert")


def test_docx_inline_math_builds_omml():
    """builder 行内 $...$ 出 m:oMath 而非字面源码（ddpm-10 实案形态）。"""
    from patent_studio.kernel.builder import PatentBuilder

    b = PatentBuilder()
    p = b.doc.add_paragraph()
    b._emit_math_text(p, r"步骤计算 $x_t=\sqrt{\alpha_t}x_{t-1}$ 后输出", None, 21, False)
    xml = p._p.xml
    assert "oMath" in xml
    assert "\\" not in xml and "sqrt" not in xml
    assert "rad" in xml  # 根式真实排版


def test_cnipa_xml_math_flatten():
    """CNIPA XML 不允许 OMML：$...$ 降级为可读线性文本（latex_to_text）。"""
    from lxml import etree as _e

    from patent_studio.kernel.cnipa_xml import inline_text

    parent = _e.fromstring("<p/>")
    inline_text(parent, r"转移核为 $p_\theta(x_{t-1}|x_t)=\mathcal{N}(\mu_\theta)$ 的核")
    out = "".join(parent.itertext())
    assert "\\" not in out and "$" not in out
    assert "θ" in out and "𝒩" in out


def test_docx_roundtrip_keeps_subsup(tmp_path):
    """sSubSup 回读无损：b_i^2 → b_i^2（docx_import 线性化可再解析）。"""
    from patent_studio.kernel.docx_import import import_docx

    d = Document()
    d.add_paragraph("一种示例装置", style="Title")
    d.add_paragraph("说明书", style="Heading 1")
    d.add_paragraph("具体实施方式", style="Heading 2")
    p = d.add_paragraph("行内公式 ")
    append_omml(p, "a_i = b_i^2")
    buf = io.BytesIO()
    d.save(buf)
    md, _rep = import_docx(buf.getvalue(), str(tmp_path))
    assert "a_i=b_i^2" in md.replace(" ", "")


def test_display_math_no_latex_residue():
    r"""端到端：ddpm-10 形态（\mathcal{N}、\left\right、\tag）经 append_omml 无残渣。"""
    d = Document()
    p = d.add_paragraph()
    append_omml(p, r"q(x_t|x_{t-1})=\mathcal{N}\left(\sqrt{1-\beta_t}x_{t-1},\beta_t I\right)\tag{1}")
    xml = p._p.xml
    assert "\\" not in xml and "mathcal" not in xml and "tag" not in xml
    assert "oMath" in xml and "rad" in xml
