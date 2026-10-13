# -*- coding: utf-8 -*-
"""单元 E3 · OA 答复工作台（docs/plan/09 §3.4）：确定性拆条 / CriticMarkup 生成律 / answer job（mock）/
意见陈述书 DOCX。"""
import importlib
import io
import os
import shutil
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.oa import answer, export, parse
from app.routes import oa as routes_oa
from patent_studio.kernel import critic
from patent_studio.kernel.ast import parse_ast

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

NOTICE = """国家知识产权局
第一次审查意见通知书
申请号：2024101234567

一、审查意见
1. 权利要求1相对于对比文件1不具备专利法第22条第2款规定的新颖性。对比文件1公开了……
其中步骤 2.5 与本申请一致。
2. 权利要求2-4、6相对于对比文件1和2的结合不具备专利法第22条第3款规定的创造性，理由如下：
（1）区别技术特征是……
（2）该区别特征是本领域的公知常识。
3. 权利要求5中「所述模块」缺乏引用基础，导致权利要求保护范围不清楚，不符合专利法第26条第4款的规定。
"""

ABOUT = """审查员意见如下。
关于权利要求1：不具备新颖性。
关于权利要求2至3：不具备创造性，容易想到。
关于说明书：公开不充分，不符合第26条第3款。
"""


# ---------------- 拆条 ----------------

def test_parse_numbered_items_claims_kinds():
    items, warnings = parse.parse(NOTICE)
    assert warnings == []
    assert [it["id"] for it in items] == ["oa-1", "oa-2", "oa-3"]
    assert items[0]["claims"] == [1] and items[0]["kind"] == "novelty"
    assert items[1]["claims"] == [2, 3, 4, 6] and items[1]["kind"] == "inventiveness"
    assert items[2]["claims"] == [5] and items[2]["kind"] == "clarity"
    # 抬头不进条目；条内的「2.5」「（1）」不拆
    assert "申请号" not in items[0]["text"] and "步骤 2.5" in items[0]["text"]
    assert "（1）" in items[1]["text"] and "（2）" in items[1]["text"]
    assert items[2]["text"].startswith("3. 权利要求5")


def test_parse_about_claims_headings():
    items, warnings = parse.parse(ABOUT)
    assert warnings == [] and len(items) == 3
    assert items[0]["claims"] == [1] and items[1]["claims"] == [2, 3]
    assert items[2]["claims"] == [] and items[2]["kind"] == "support"
    assert items[0]["text"] == "关于权利要求1：不具备新颖性。"


def test_parse_fallback_whole_text_and_empty():
    items, warnings = parse.parse("审查员认为权利要求 1 和 3 不清楚。")
    assert len(items) == 1 and items[0]["id"] == "oa-1" and items[0]["claims"] == [1, 3] and warnings
    assert parse.parse("  \n") == ([], ["通知书文本为空"])


def test_parse_is_deterministic():
    assert parse.parse(NOTICE) == parse.parse(NOTICE)


# ---------------- CriticMarkup 生成律 ----------------

@pytest.mark.parametrize("old,new", [
    ("1. 一种方法，其特征在于，包括步骤A。", "1. 一种方法，其特征在于，包括步骤A和步骤B。"),
    ("1. 一种方法，包括：\n步骤一；\n步骤二。", "1. 一种方法，其特征在于，包括：\n步骤一；"),
    ("1. 甲。", "1. 甲。\n乙。\n丙。"),
    ("1. 完全不同的原文。", "1. 全新的替换文本，与原文毫无关系。"),
])
def test_to_critic_roundtrip(old, new):
    line = answer.to_critic(old, new)
    assert critic.reject(line).rstrip("\n") == old
    assert critic.accept(line).rstrip("\n") == new  # 整行删除处 accept 后留空行
    assert critic.lint_marks(line) == []
    assert critic.has_markup(line)


def test_to_critic_merges_fragments():
    line = answer.to_critic("所述模块连接至所述单元", "所述装置连接至所述部件")
    assert line.count("{~~") <= 2 and critic.reject(line) == "所述模块连接至所述单元"
    assert critic.accept(line) == "所述装置连接至所述部件"


# ---------------- 端到端（mock 供应商）----------------

@pytest.fixture()
def client(tmp_path, monkeypatch):
    projects = tmp_path / "projects"
    projects.mkdir()
    shutil.copytree(os.path.join(ROOT, "examples", "method"), projects / "method",
                    ignore=shutil.ignore_patterns("*.docx", ".studio", "*.bak-*"))
    monkeypatch.setenv("PATENT_PROJECTS_DIR", str(projects))
    monkeypatch.setenv("PATENT_STUDIO_AI_PROVIDER", "mock")
    from app import config, main
    from app.domain import build, xml_export
    from app.domain import projects as projects_mod

    for m in (config, projects_mod, build, xml_export, main):
        importlib.reload(m)
    with TestClient(main.app) as c:
        yield c, projects, main.app


def _wait(c, job_id, timeout=10.0):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = c.get(f"/api/ai/jobs/{job_id}").json()
        if j["status"] in ("done", "error"):
            return j
        time.sleep(0.05)
    raise AssertionError(f"job {job_id} 未在 {timeout}s 内结束：{j}")


def test_oa_parse_route_persists_items(client):
    c, projects, app = client
    r = c.post("/api/projects/method/oa/parse", json={"text": NOTICE})
    assert r.status_code == 200
    body = r.json()
    assert [it["id"] for it in body["items"]] == ["oa-1", "oa-2", "oa-3"] and body["warnings"] == []
    assert (projects / "method" / ".studio" / "oa" / "items.json").exists()
    assert app.state.ai.provider.calls == []  # 拆条不叫 AI
    assert c.post("/api/projects/nope/oa/parse", json={"text": NOTICE}).status_code == 404


def test_oa_answer_requires_ai_enabled_and_items(client):
    c, _, app = client
    r = c.post("/api/projects/method/oa/answer", json={})
    assert r.status_code == 403 and r.json()["detail"]["kind"] == "ai_disabled"
    c.put("/api/projects/method/settings", json={"ai_enabled": True})
    r = c.post("/api/projects/method/oa/answer", json={})
    assert r.status_code == 409 and r.json()["detail"]["kind"] == "oa_empty"
    assert app.state.ai.provider.calls == []
    assert c.get("/api/projects/method/oa/export.docx").status_code == 409


def test_oa_answer_job_and_export_docx(client):
    c, projects, app = client
    pdir = projects / "method"
    before = (pdir / "content.md").read_text(encoding="utf-8")
    c.post("/api/projects/method/oa/parse", json={"text": NOTICE})
    c.put("/api/projects/method/settings", json={"ai_enabled": True})

    r = c.post("/api/projects/method/oa/answer", json={"item_ids": ["oa-1", "oa-3"]})
    assert r.status_code == 202, r.text
    job = r.json()
    assert job["kind"] == "oa_answer" and job["pid"] == "method"
    job = _wait(c, job["id"])
    assert job["status"] == "done", job
    answers = job["result"]["answers"]
    assert [a["item_id"] for a in answers] == ["oa-1", "oa-3"]
    assert [x["kind"] for x in app.state.ai.provider.calls] == ["oa_answer", "oa_answer"]

    a = answers[0]
    assert a["points"] and a["amend_summary"]
    assert len(a["claim_edits"]) == 1
    e = a["claim_edits"][0]
    assert e["claim"] == 1 and e["block_id"] == "claims/1" and critic.has_markup(e["critic_line"])
    assert critic.lint_marks(e["critic_line"]) == []
    # 建议可还原原权项；accept 后含 mock 增补的特征
    content = (pdir / "content.md").read_text(encoding="utf-8")
    assert content == before  # 不动工作区
    blocks = answer.claim_blocks(content)
    assert critic.reject(e["critic_line"]) == blocks[1]["text"] == e["original"]
    assert "特征 A" in critic.accept(e["critic_line"])
    assert e["line"] == blocks[1]["line"] and content.split("\n")[e["line"] - 1].startswith("1.")
    # 经 patch.py 写回后仍可解析，且 reject 回原文
    patched = c.post("/api/projects/method/blocks/claims/1", json={"markdown": e["critic_line"]})
    assert patched.status_code == 200, patched.text
    new_content = (pdir / "content.md").read_text(encoding="utf-8")
    assert critic.reject(new_content) == before and parse_ast(new_content)["claims"][0]["num"] == 1
    assert sorted(os.listdir(pdir / ".studio" / "oa")) == ["items.json", "oa-1.json", "oa-3.json"]

    # 导出意见陈述书：python-docx 可打开，答复表含每条编号
    r = c.get("/api/projects/method/oa/export.docx")
    assert r.status_code == 200 and r.headers["content-type"].startswith(routes_oa.DOCX_MIME)
    d = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "意见陈述书" in text and "修改后的权利要求书" in text
    assert len(d.tables) == 1 and len(d.tables[0].rows) == 4  # 表头 + 3 条
    col0 = [row.cells[0].text for row in d.tables[0].rows]
    assert col0 == ["序号", "1", "2", "3"]
    assert "（尚未答复）" in d.tables[0].rows[2].cells[2].text  # oa-2 未生成答复
    assert a["points"][0] in d.tables[0].rows[1].cells[2].text

    # 用户改要点 → PUT 落盘 → 导出以改后要点为准
    r = c.put("/api/projects/method/oa/answers/oa-1", json={"points": ["  改后要点甲 ", "", "改后要点乙"]})
    assert r.status_code == 200 and r.json()["points"] == ["改后要点甲", "改后要点乙"]
    assert r.json()["claim_edits"] == a["claim_edits"]
    assert c.put("/api/projects/method/oa/answers/oa-2", json={"points": ["x"]}).status_code == 404
    d = Document(io.BytesIO(c.get("/api/projects/method/oa/export.docx").content))
    cell = d.tables[0].rows[1].cells[2].text
    assert "改后要点乙" in cell and a["points"][0] not in cell


def test_export_docx_amended_claims_uses_accept():
    claims_md = "\n1. 一种方法，其特征在于，包括步骤A。\n\n2. 根据权利要求1所述的方法。\n\n"
    ans = {"item_id": "oa-1", "points": ["要点"], "amend_summary": "增补 B",
           "claim_edits": [{"claim": 1, "block_id": "claims/1", "line": 1,
                            "critic_line": "1. 一种方法，其特征在于，包括步骤A{++和步骤B++}。",
                            "original": "1. 一种方法，其特征在于，包括步骤A。"}]}
    items = [{"id": "oa-1", "claims": [1], "kind": "novelty", "text": "1. 无新颖性"}]
    data = export.export_docx(project={"title": "T", "app_no": None}, items=items, answers=[ans], claims_md=claims_md)
    d = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "1. 一种方法，其特征在于，包括步骤A和步骤B。" in text and "{++" not in text
    assert "2. 根据权利要求1所述的方法。" in text


def test_oa_items_roundtrip_restores_workset(client):
    """D-OA-1：拆条结果+答复落盘后，GET /oa/items 须能重建工作集（刷新不丢条目）。"""
    c, projects, _ = client
    r = c.get("/api/projects/method/oa/items")
    assert r.status_code == 200 and r.json() == {"items": [], "warnings": [], "answers": []}

    c.post("/api/projects/method/oa/parse", json={"text": NOTICE})
    c.put("/api/projects/method/settings", json={"ai_enabled": True})
    job = c.post("/api/projects/method/oa/answer", json={"item_ids": ["oa-1"]}).json()
    _wait(c, job["id"])

    r = c.get("/api/projects/method/oa/items")
    assert r.status_code == 200
    body = r.json()
    assert [it["id"] for it in body["items"]] == ["oa-1", "oa-2", "oa-3"]
    assert [a["item_id"] for a in body["answers"]] == ["oa-1"]
    assert body["answers"][0]["points"]  # 答复要点随工作集返回

    # 用户改要点 → PUT 落盘 → 重建读到新版本
    c.put("/api/projects/method/oa/answers/oa-1", json={"points": ["改过的要点"]})
    body = c.get("/api/projects/method/oa/items").json()
    assert body["answers"][0]["points"] == ["改过的要点"]

    assert c.get("/api/projects/nope/oa/items").status_code == 404
