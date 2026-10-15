# -*- coding: utf-8 -*-
"""单元 E3 · 一键转专利（docs/plan/09 §3.3）：抽取 / 问卷 / 渲染去号 / mock 端到端（新案卷 → 三书装配律 → build；
现案卷 → history 草稿 commit 不动工作区）/ 校验失败回修 1 次 / ai_ack 与 ai_enabled 门。"""

import importlib
import io
import json
import os
import shutil
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app.ai.client import Usage
from app.ai.mock import MockProvider
from app.convert import extract, pipeline, questionnaire
from app.domain.projects import ProjectStore
from app.routes.convert import new_pid
from patent_studio.kernel import docset, figures_md
from patent_studio.kernel.ast import parse_ast

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIXTURE = json.load(
    open(os.path.join(os.path.dirname(__file__), "..", "app", "ai", "fixtures", "convert.json"), encoding="utf-8")
)
PAPER = (
    "A method for adaptive scheduling of heterogeneous computing tasks.\n\n"
    "We propose a queue based scheduler that measures task latency and\n"
    "re-assigns tasks between CPU and accelerator pools.\n"
) * 3
PAPER_BYTES = PAPER.encode("utf-8")


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(buf, format="PNG")
    return buf.getvalue()


def _docx() -> bytes:
    d = Document()
    d.add_paragraph("第一段：一种调度方法。")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "参数", "取值"
    d.add_picture(io.BytesIO(_png()))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(text: str | None) -> bytes:
    """单页 PDF，Helvetica 直接写 ASCII；text=None ⇒ 无文字（模拟扫描件）。"""
    stream = b"" if text is None else b"BT /F1 12 Tf 72 720 Td (" + text.encode("latin-1") + b") Tj ET"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


# ---------------- 抽取 ----------------


def test_extract_txt_md_normalizes():
    r = extract.extract("a.txt", "第一行\r\n\r\n\r\n\r\n第二行\r\n".encode())
    assert r["text"] == "第一行\n\n第二行" and r["pages"] == 0 and r["figures"] == [] and r["warnings"] == []
    assert extract.extract("b.md", b"# T\n\nbody")["text"] == "# T\n\nbody"


def test_extract_docx_paragraphs_tables_figures():
    r = extract.extract("x.docx", _docx())
    assert "第一段：一种调度方法。" in r["text"] and "| 参数 | 取值 |" in r["text"]
    assert len(r["figures"]) == 1 and r["figures"][0][0] == "png" and r["figures"][0][1][:4] == b"\x89PNG"


def test_extract_pdf_text_and_scanned():
    r = extract.extract("p.pdf", _pdf("Hello adaptive scheduling of tasks"))
    assert "Hello adaptive scheduling" in r["text"] and r["pages"] == 1
    with pytest.raises(extract.ExtractError) as ei:
        extract.extract("scan.pdf", _pdf(None))
    assert ei.value.kind == "scanned_pdf" and ei.value.status == 422


def _pdf_vector_page() -> bytes:
    """一页矢量图 + 「Figure 1:」题注（无内嵌位图——TikZ 类论文形态）。"""
    stream = (
        b"0 0 1 rg 100 560 400 140 re S "
        b"0 0 0 rg 100 630 m 500 630 l S 100 570 m 500 570 l S "
        b"BT /F1 12 Tf 72 500 Td (Figure 1: A schematic view of the module) Tj ET "
        b"BT /F1 12 Tf 72 100 Td (Body text line one.) Tj ET"
    )
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


_POPPLER = pytest.mark.skipif(
    not (shutil.which("pdftotext") and shutil.which("pdftoppm")), reason="poppler-utils 不在测试环境"
)


@_POPPLER
def test_extract_pdf_vector_figure_fallback():
    """往复轮114——TikZ/矢量图形兜底：无内嵌位图的页按题注墨带渲染 png。
    mamba(2312.00752) 实案：全部图为 PGF 矢量，位图通道 0 产出。"""
    data = _pdf_vector_page()
    figs, thumbs, dropped, academic, caps = extract._pdf_vector_figures(data, 5)
    assert dropped == 0 and len(figs) == 1 and figs[0][0] == "png"
    assert thumbs  # 墨带指纹供位图通道去重
    assert len(caps) == len(figs) and caps[0] and "Figure" in caps[0]
    from PIL import Image

    im = Image.open(io.BytesIO(figs[0][1]))
    assert im.width >= 400 and im.height >= 150  # 裁到矩形墨带而非碎片
    r = extract.extract("v.pdf", data)
    assert len(r["figures"]) == 1
    assert any("题注" in w for w in r["warnings"])


def test_figure_ok_drops_strips_blank_undecodable():
    """往复轮105 抽图质量闸——YOLO fig2 7013×476 白粉条实案：条形/近空白/坏图剔除，正常图与 emf 放行。"""
    from PIL import Image

    def _img(w: int, h: int, rgb: tuple) -> bytes:
        buf = io.BytesIO()
        Image.new("RGB", (w, h), rgb).save(buf, "PNG")
        return buf.getvalue()

    assert extract._figure_ok(_img(64, 64, (30, 30, 30))) is True
    assert extract._figure_ok(_img(7013, 476, (255, 255, 255))) is False  # 条形碎片
    assert extract._figure_ok(_img(200, 200, (255, 255, 255))) is False  # 近空白
    assert extract._figure_ok(b"not-an-image") is False  # 无法解码
    assert extract._figure_ok(b"\x01\x00\x00\x00emf-body") is True  # emf 矢量放行


def test_figure_ok_flattens_alpha_and_normalize_grayscales():
    """Transformer 实案：cv-fig2-5 是 RGB 全黑、alpha 近全 0 的透明空壳——白底合成后近空白剔除；
    cv-fig1 彩色雷达图经 _normalize_fig 白底合成+转灰度存 PNG（CNIPA 附图禁彩色）。"""
    from PIL import Image

    from app.convert.pipeline import _normalize_fig

    def _rgba_shell(w: int, h: int) -> bytes:
        buf = io.BytesIO()
        im = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        # 四周圆框细线留少量不透明像素——复刻实案「空框假象」
        for x in range(w):
            im.putpixel((x, 0), (0, 0, 0, 255))
            im.putpixel((x, h - 1), (0, 0, 0, 255))
        for y in range(h):
            im.putpixel((0, y), (0, 0, 0, 255))
            im.putpixel((w - 1, y), (0, 0, 0, 255))
        im.save(buf, "PNG")
        return buf.getvalue()

    shell = _rgba_shell(1470, 723)
    assert extract._figure_ok(shell) is False  # 合成白底后近空白 → 剔

    def _color_img() -> bytes:
        buf = io.BytesIO()
        Image.new("RGB", (64, 64), (200, 40, 40)).save(buf, "PNG")
        return buf.getvalue()

    ext, blob = _normalize_fig("png", _color_img())
    assert ext == "png"
    im = Image.open(io.BytesIO(blob))
    assert im.mode == "L"  # 彩色 → 灰度
    assert _normalize_fig("emf", b"\x01\x00\x00\x00x") == ("emf", b"\x01\x00\x00\x00x")  # 矢量原样放行


def test_extract_rejects_unsupported_and_oversize():
    with pytest.raises(extract.ExtractError) as ei:
        extract.extract("a.pptx", b"x")
    assert ei.value.kind == "unsupported_media" and ei.value.status == 415
    with pytest.raises(extract.ExtractError) as ei:
        extract.extract("a.txt", b"x" * (extract.MAX_BYTES + 1))
    assert ei.value.kind == "too_large" and ei.value.status == 413


def test_extract_caps_chars():
    r = extract.extract("a.txt", ("字" * (extract.MAX_CHARS + 100)).encode("utf-8"))
    assert len(r["text"]) == extract.MAX_CHARS and any("截断" in w for w in r["warnings"])


# ---------------- 问卷 ----------------


def test_questionnaire_parse_and_render():
    q = questionnaire.parse(json.dumps({"problem": "延迟高", "effects": "  ", "unknown": "x", "figures": "图1 流程"}))
    assert q == {"problem": "延迟高", "figures": "图1 流程"}
    assert questionnaire.parse(None) == {} and questionnaire.parse("") == {}
    text = questionnaire.render(q)
    assert text.startswith("要解决的技术问题：\n延迟高") and "附图说明：\n图1 流程" in text
    for bad in ("[1]", "{", "null"):
        with pytest.raises(questionnaire.QuestionnaireError):
            questionnaire.parse(bad)


# ---------------- 渲染：编号由内核生成 ----------------


def test_render_markdown_strips_ai_numbering_and_is_valid():
    data = {
        "title": "一种调度方法",
        "abstract": "摘要一段。",
        "claims": [
            "1. 一种调度方法，其特征在于，包括步骤 A。",
            "2、根据权利要求1所述的方法，其特征在于，还包括步骤 B。",
        ],
        "description": {
            "技术领域": "[0001] 本发明涉及计算机领域。",
            "背景技术": ["### 背景技术", "[0002] 现有技术……"],
            "发明内容": "1. 本发明提供一种方法。\n\n2. 有益效果如下。",
            "附图说明": "图1 为流程图。",
            "具体实施方式": "0005 实施例一。",
        },
    }
    md = pipeline.render_markdown(data)
    assert md.startswith("# 一种调度方法\n\n## 摘要\n\n摘要一段。\n\n## 权利要求书\n\n1.一种调度方法")
    assert "\n2.根据权利要求1所述的方法" in md and "2、" not in md
    assert "[0001]" not in md and "0005 " not in md and md.count("### 背景技术") == 1
    assert "\n### 发明内容\n\n1. 本发明提供一种方法。\n\n2. 有益效果如下。\n" in md
    assert pipeline.check(md) == []
    assert docset.assemble(docset.split(md)) == md
    ast = parse_ast(md)
    assert [c["num"] for c in ast["claims"]] == [1, 2]


def test_fixture_renders_valid_docset():
    md = pipeline.render_markdown(FIXTURE)
    assert pipeline.check(md) == [] and docset.assemble(docset.split(md)) == md


def test_render_markdown_caps_overlong_abstract():
    """往复轮107a：摘要超限确定性截断——sam2 实案 AI 摘要 309 字两次回修仍超 300，
    整案转换报废（ai_format）。现与 s1 共用 preflight.limit_abstract 标点边界截断。"""
    data = dict(FIXTURE, abstract="本发明公开一种测试方法，解决了调度问题。" * 20)
    warnings: list[str] = []
    md = pipeline.render_markdown(data, warnings)
    assert pipeline.check(md) == []
    assert any("截断" in w for w in warnings)
    import re as _re

    abs_sec = _re.search(r"## 摘要\n\n(.*?)\n\n## 权利要求书", md, _re.S).group(1)
    assert len(abs_sec.replace("\n", "")) <= 290


def test_render_markdown_normalizes_self_reference():
    """往复轮108：convert 草稿自称「本申请」——s1 早已归一本发明（轮77），
    convert 链未接同款处理，sam2 实案草稿 13 处本申请贯穿三书。"""
    data = dict(
        FIXTURE,
        abstract="本申请公开一种测试方法。",
        claims=["本申请方法，其特征在于，包括步骤。"],
        description=dict(FIXTURE["description"], 技术领域=["本申请涉及测试领域。"]),
    )
    md = pipeline.render_markdown(data)
    assert "本申请" not in md
    assert md.count("本发明") >= 3


def test_convert_retries_transient_ai_error(monkeypatch, tmp_path):
    """往复轮107b：retryable AiError（超时/429/5xx/断连）同批再试一次——
    mamba 实案主模型与兜底模型接连 ReadTimeout 致整案报废。"""
    import asyncio
    import types

    from app.ai.client import AiError

    calls = {"n": 0}

    async def flaky(ai, messages, *, project_dir, job_id):
        calls["n"] += 1
        if calls["n"] == 1:
            raise AiError("timeout", "AI 中转站超时：ReadTimeout", 504, retryable=True)
        return (json.dumps(FIXTURE, ensure_ascii=False), {})

    async def no_sleep(_s):
        return None

    monkeypatch.setattr(pipeline.runner, "heavy_chat", flaky)
    monkeypatch.setattr(pipeline.asyncio, "sleep", no_sleep)

    store = types.SimpleNamespace(
        _dir=lambda pid: str(tmp_path / pid),
        save_content=lambda pid, content: (tmp_path / pid / "content.md").write_text(content, encoding="utf-8"),
        encoding="utf-8",
    )
    extracted = {"text": "paper", "warnings": [], "figures": [], "pages": 1}
    q = questionnaire.parse("{}")
    (tmp_path / "case1").mkdir()
    result = asyncio.run(
        pipeline.run(
            store=store,
            ai=None,
            pid="case1",
            is_new=True,
            filename="p.txt",
            extracted=extracted,
            questionnaire=q,
            progress=lambda *a: None,
        )
    )
    assert calls["n"] == 2
    assert result["report"]["claims"] >= 1
    assert (tmp_path / "case1" / "content.md").exists()


def test_gen_quality_errors_flags_abstract_and_marks():
    """生成质量回修门（patchtst 实案）：摘要缺细则 23 条要素 / figures.marks 声明的
    标记号未进正文——与 preflight abstract.elements_missing / numerals.mark_not_in_spec
    同口径，convert 首轮按 error 回修。"""
    assert pipeline._gen_quality_errors(FIXTURE) == []
    bad_abs = dict(FIXTURE, abstract="本发明公开一种调度方法，属于调度领域，包括步骤 A，能够降低延迟。")
    errs = pipeline._gen_quality_errors(bad_abs)
    assert any("摘要未体现" in e and "技术问题" in e for e in errs)
    marked = dict(FIXTURE, figures=[{"num": 2, "caption": "数据流向图", "marks": {"207": "通道预测拼接"}}])
    errs = pipeline._gen_quality_errors(marked)
    assert any("207" in e and "正文" in e for e in errs)
    ok = dict(marked, description=dict(FIXTURE["description"], 具体实施方式="各通道输出经通道预测拼接207合并。"))
    assert pipeline._gen_quality_errors(ok) == []
    # 数字边界：正文里 2077/1207 不算引用了 207
    near = dict(marked, description=dict(FIXTURE["description"], 具体实施方式="第2077个参数与1207无关。"))
    assert any("207" in e for e in pipeline._gen_quality_errors(near))


def _stub_store(tmp_path):
    import types

    return types.SimpleNamespace(
        _dir=lambda pid: str(tmp_path / pid),
        save_content=lambda pid, content: (tmp_path / pid / "content.md").write_text(content, encoding="utf-8"),
    )


def _run_convert(store, tmp_path, pid="case1"):
    import asyncio

    extracted = {"text": "paper", "warnings": [], "figures": [], "pages": 1}
    (tmp_path / pid).mkdir(exist_ok=True)
    return asyncio.run(
        pipeline.run(
            store=store,
            ai=None,
            pid=pid,
            is_new=True,
            filename="p.txt",
            extracted=extracted,
            questionnaire=questionnaire.parse("{}"),
            progress=lambda *a: None,
        )
    )


def test_convert_regenerates_on_gen_quality(monkeypatch, tmp_path):
    """首轮质量问题 → 带错误回修，第二轮合格即交付。"""
    import json as _json

    bad = dict(FIXTURE, abstract="本发明公开一种调度方法，属于调度领域，包括步骤 A，能够降低延迟。")
    calls = {"n": 0}

    async def gen(ai, messages, *, project_dir, job_id):
        calls["n"] += 1
        return (_json.dumps(bad if calls["n"] == 1 else FIXTURE, ensure_ascii=False), {})

    monkeypatch.setattr(pipeline.runner, "heavy_chat", gen)
    result = _run_convert(_stub_store(tmp_path), tmp_path)
    assert calls["n"] == 2 and result["report"]["claims"] >= 1
    assert not any("摘要未体现" in w for w in result["report"]["warnings"])


def test_convert_gen_quality_degrades_to_warning(monkeypatch, tmp_path):
    """回修一轮后质量问题仍在 → 不报废整案，退为告警交付（与 preflight warning 同口径）。"""
    import json as _json

    bad = dict(FIXTURE, abstract="本发明公开一种调度方法，属于调度领域，包括步骤 A，能够降低延迟。")
    calls = {"n": 0}

    async def gen(ai, messages, *, project_dir, job_id):
        calls["n"] += 1
        return (_json.dumps(bad, ensure_ascii=False), {})

    monkeypatch.setattr(pipeline.runner, "heavy_chat", gen)
    result = _run_convert(_stub_store(tmp_path), tmp_path)
    assert calls["n"] == 2
    assert any("摘要未体现" in w for w in result["report"]["warnings"])
    assert (tmp_path / "case1" / "content.md").exists()


# ---------------- 端到端（mock 供应商）----------------


@pytest.fixture()
def client(tmp_path, monkeypatch):
    projects = tmp_path / "projects"
    projects.mkdir()
    shutil.copytree(
        os.path.join(ROOT, "examples", "method"),
        projects / "method",
        ignore=shutil.ignore_patterns("*.docx", ".studio", "*.bak-*"),
    )
    monkeypatch.setenv("PATENT_PROJECTS_DIR", str(projects))
    monkeypatch.setenv("PATENT_STUDIO_AI_PROVIDER", "mock")
    from app import config, main
    from app.domain import build, xml_export
    from app.domain import projects as projects_mod

    for m in (config, projects_mod, build, xml_export, main):
        importlib.reload(m)
    with TestClient(main.app) as c:
        yield c, projects, main.app


def _wait(c, job_id, timeout=10.0):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = c.get(f"/api/ai/jobs/{job_id}").json()
        if j["status"] in ("done", "error"):
            return j
        time.sleep(0.05)
    raise AssertionError(f"job {job_id} 未在 {timeout}s 内结束：{j}")


def _convert(c, target="new", ai_ack=True, name="scheduler paper.txt", data=None, q="{}"):
    data = PAPER_BYTES if data is None else data
    return c.post(
        "/api/projects/convert",
        files={"file": (name, data, "text/plain")},
        data={"questionnaire": q, "target": target, "ai_ack": "true" if ai_ack else "false"},
    )


def test_convert_new_project_end_to_end(client):
    c, projects, app = client
    r = _convert(c, q=json.dumps({"problem": "调度延迟高"}))
    assert r.status_code == 202, r.text
    job = r.json()
    assert job["kind"] == "convert" and job["status"] in ("queued", "running") and job["pid"] == "scheduler-paper"
    job = _wait(c, job["id"])
    assert job["status"] == "done", job
    res = job["result"]
    assert res["new_pid"] == "scheduler-paper" and res["draft_sha"] is None
    rep = res["report"]
    assert rep["claims"] == len(FIXTURE["claims"]) and rep["chars_in"] == len(PAPER.strip())
    assert rep["usage"]["total_tokens"] > 0
    assert rep["lint"]["error"] == 0 and rep["notes"] == FIXTURE["notes"]

    # mock 只被叫了一次，且提示词是 convert，问卷进了提示词
    prov = app.state.ai.provider
    assert isinstance(prov, MockProvider) and [x["kind"] for x in prov.calls] == ["convert"]
    assert prov.calls[0]["model"] == "swe-2-max" and prov.calls[0]["stream"] is False

    # 新案卷：三书装配律 + parse_ast + build 200
    pdir = projects / "scheduler-paper"
    # 文本源无提取图：占位图须被清空且 figures.md 回写空 spec，不留悬空声明
    assert not list((pdir / "figs").iterdir()) and (pdir / "docs" / "claims.md").exists()
    assert figures_md.read(str(pdir)).spec == []
    content = c.get("/api/projects/scheduler-paper/content").json()["content"]
    assert docset.check([str(pdir)]) == []
    assert docset.assemble(docset.split(content)) == content
    assert len(parse_ast(content)["claims"]) == rep["claims"]
    b = c.post("/api/projects/scheduler-paper/build", json={"pdf": False})
    assert b.status_code == 200 and b.json()["status"] == "ok", b.text
    assert (pdir / ".studio" / "ai-usage.jsonl").read_text(encoding="utf-8").count("\n") == 1
    assert c.get("/api/projects/scheduler-paper/ai/jobs").json()["jobs"][0]["id"] == job["id"]


def test_convert_new_requires_ack_and_makes_no_call(client):
    c, projects, app = client
    r = _convert(c, ai_ack=False)
    assert r.status_code == 403 and r.json()["detail"]["kind"] == "ai_disabled"
    assert app.state.ai.provider.calls == [] and set(os.listdir(projects)) == {"method"}


def test_convert_from_attachment(client):
    """输入二选一：attachment_pid + attachment 直接读 <project>/attachments/<name>（W4-C）。"""
    c, projects, app = client
    (projects / "method" / "attachments").mkdir()
    (projects / "method" / "attachments" / "scheduler paper.txt").write_bytes(PAPER_BYTES)
    base = {"questionnaire": "{}", "target": "new", "ai_ack": "true"}
    r = c.post("/api/projects/convert", data=base)
    assert r.status_code == 422 and r.json()["detail"]["kind"] == "source"
    r = c.post(
        "/api/projects/convert",
        files={"file": ("a.txt", PAPER_BYTES, "text/plain")},
        data={**base, "attachment_pid": "method", "attachment": "scheduler paper.txt"},
    )
    assert r.status_code == 422 and r.json()["detail"]["kind"] == "source"
    r = c.post("/api/projects/convert", data={**base, "attachment_pid": "method", "attachment": "nope.txt"})
    assert r.status_code == 404 and r.json()["detail"]["kind"] == "attachments"
    r = c.post("/api/projects/convert", data={**base, "attachment_pid": "ghost", "attachment": "scheduler paper.txt"})
    assert r.status_code == 404 and "ghost" in r.json()["detail"]
    r = c.post("/api/projects/convert", data={**base, "attachment_pid": "method", "attachment": "scheduler paper.txt"})
    assert r.status_code == 202, r.text
    assert r.json()["pid"] == "scheduler-paper"
    job = _wait(c, r.json()["id"])
    assert job["status"] == "done" and job["result"]["report"]["chars_in"] == len(PAPER.strip())
    assert (projects / "method" / "attachments" / "scheduler paper.txt").read_bytes() == PAPER_BYTES


def test_convert_pid_dedupe_and_slug(tmp_path):
    store = ProjectStore(str(tmp_path))
    assert new_pid(store, "我的 论文(v2).pdf") == "我的-论文-v2"
    assert new_pid(store, "...") == "convert"


def test_convert_new_pid_collision(client):
    c, projects, _ = client
    r = _convert(c, name="method.txt")
    assert r.status_code == 202 and r.json()["pid"] == "method-2"
    assert _wait(c, r.json()["id"])["status"] == "done"


def test_convert_existing_project_draft_commit(client):
    c, projects, app = client
    r = _convert(c, target="method")
    assert r.status_code == 403 and r.json()["detail"]["kind"] == "ai_disabled"
    assert app.state.ai.provider.calls == []

    c.put("/api/projects/method/settings", json={"ai_enabled": True})
    before = (projects / "method" / "content.md").read_text(encoding="utf-8")
    r = _convert(c, target="method")
    assert r.status_code == 202, r.text
    job = _wait(c, r.json()["id"])
    assert job["status"] == "done", job
    res = job["result"]
    assert res["new_pid"] is None and res["draft_sha"] and res["pid"] == "method"
    assert (projects / "method" / "content.md").read_text(encoding="utf-8") == before
    assert c.get("/api/projects/method/content").json()["content"] == before
    from app.domain import history

    pdir = str(projects / "method")
    top = history.log(pdir)[0]
    assert top["sha"] == res["draft_sha"] and top["message"] == pipeline.DRAFT_MESSAGE and top["author"] == "ai"
    draft = history.show(pdir, res["draft_sha"])
    assert draft != before and pipeline.check(draft) == []
    assert (projects / "method" / "figs" / "fig1.png").exists()


def test_convert_existing_project_stages_figures_as_attachments(client):
    """草稿模式抽出的图不落 figs/，写 attachments/cv-figN 供采纳后晋级。"""
    c, projects, _ = client
    c.put("/api/projects/method/settings", json={"ai_enabled": True})
    figs_dir = projects / "method" / "figs"
    before_figs = set(os.listdir(figs_dir))
    r = _convert(c, target="method", name="x.docx", data=_docx())
    assert r.status_code == 202, r.text
    job = _wait(c, r.json()["id"])
    assert job["status"] == "done", job
    rep = job["result"]["report"]
    staged = projects / "method" / "attachments" / "cv-fig1.png"
    assert staged.exists() and staged.read_bytes()[:4] == b"\x89PNG"
    assert set(os.listdir(figs_dir)) == before_figs
    assert any("附件" in w for w in rep["warnings"])
    # 二次转换同名不覆盖——cv-fig1-2.png
    r = _convert(c, target="method", name="y.docx", data=_docx())
    assert _wait(c, r.json()["id"])["status"] == "done"
    assert (projects / "method" / "attachments" / "cv-fig1-2.png").exists()


def test_convert_busy_409_and_atomic_create(client):
    """R75-01：已有进行中的 convert job ⇒ 409；create_unless_busy 同步原子（并发双过闸的兜底）。"""
    c, _, app = client
    c.put("/api/projects/method/settings", json={"ai_enabled": True})
    jobs = app.state.jobs
    live = jobs.create("convert", "method")  # queued ⇒ busy
    assert jobs.busy("method", "convert")
    r = _convert(c, target="method")
    assert r.status_code == 409 and r.json()["detail"]["kind"] == "job_busy", r.text
    # 原子闸门：busy 时直接返回 None（不再产出第二个 job）
    assert jobs.create_unless_busy("convert", "method") is None
    assert len([j for j in jobs.list("method") if j.kind == "convert"]) == 1
    # 无关 kind / 其他 pid 不受 busy 影响
    assert jobs.create_unless_busy("review", "method") is not None
    assert jobs.create_unless_busy("convert", "system") is not None
    # job 终态后闸门放开
    jobs.update(live, status="done")
    assert jobs.create_unless_busy("convert", "method") is not None


def test_convert_docx_figures_written(client):
    c, projects, _ = client
    r = _convert(c, name="disclosure.docx", data=_docx())
    assert r.status_code == 202, r.text
    job = _wait(c, r.json()["id"])
    assert job["status"] == "done", job
    assert job["result"]["report"]["figures"] == 1
    assert sorted(os.listdir(projects / "disclosure" / "figs")) == ["fig1.png"]
    assert (projects / "disclosure" / "figs" / "fig1.png").read_bytes()[:4] == b"\x89PNG"


def test_write_figures_resyncs_figures_md(tmp_path):
    """非 PNG 首图（PDF 抽出的 tiff 等）落 figs/ 后 figures.md 仍指占位 fig1.png——
    必须按真实文件回写，否则 spec_figures/abstract_figure 悬空缺图。"""
    pdir = tmp_path / "proj"
    (pdir / "figs").mkdir(parents=True)
    (pdir / "docs").mkdir()
    (pdir / "figs" / "fig1.png").write_bytes(b"placeholder")
    (pdir / "docs" / "figures.md").write_text(
        "# 说明书附图\n\n![图1](figs/fig1.png){width=15}\n\n# 摘要附图\n\n![](figs/fig1.png){width=14.5}\n",
        encoding="utf-8",
    )
    assert (
        pipeline._write_figures(
            str(pdir), [("tiff", b"T1"), ("png", b"P2"), ("png", b"P3")], "图1为架构图\n图2为流程图\n图3为结果图"
        )
        == 3
    )
    fmd = (pdir / "docs" / "figures.md").read_text(encoding="utf-8")
    assert "fig1.png" not in fmd
    doc = figures_md.read(str(pdir))
    assert doc.abstract is not None and doc.abstract.file == "fig1.tiff"
    assert [f.file for f in doc.spec] == ["fig1.tiff", "fig2.png", "fig3.png"]


def test_write_figures_caps_spec_at_declared(tmp_path):
    """抽图多于草稿声明图号 → 超出的只留 figs/ 素材不进 spec；声明 0 图时 spec 空、摘要图保底首图。"""
    pdir = tmp_path / "proj"
    (pdir / "figs").mkdir(parents=True)
    figs = [("png", b"P1"), ("png", b"P2"), ("png", b"P3"), ("png", b"P4"), ("png", b"P5")]
    assert pipeline._write_figures(str(pdir), figs, "…图1为架构图…图2为流程图…") == 5
    doc = figures_md.read(str(pdir))
    assert [f.file for f in doc.spec] == ["fig1.png", "fig2.png"]
    assert doc.abstract is not None and doc.abstract.file == "fig1.png"
    assert len(os.listdir(pdir / "figs")) == 5  # 素材全留盘

    pdir2 = tmp_path / "proj2"
    (pdir2 / "figs").mkdir(parents=True)
    pipeline._write_figures(str(pdir2), [("png", b"P1"), ("png", b"P2")], "无图号引用的正文")
    doc2 = figures_md.read(str(pdir2))
    assert doc2.spec == [] and doc2.abstract is not None


def test_convert_scanned_pdf_422(client):
    c, _, _ = client
    r = _convert(c, name="scan.pdf", data=_pdf(None))
    assert r.status_code == 422 and r.json()["detail"]["kind"] == "scanned_pdf"


# ---------------- 校验失败回修 1 次 ----------------


class SeqProvider:
    """按序返回预设文本；用于验证「失败带错误回修 1 次」。"""

    name = "seq"
    key = "k"

    def __init__(self, texts):
        self.texts = list(texts)
        self.calls = []

    async def complete(self, model, messages, *, timeout, max_tokens=None, temperature=0.0):
        self.calls.append(messages)
        return self.texts.pop(0), Usage(10, 10, 20)


BROKEN = dict(FIXTURE, claims=[])  # 无权项 ⇒ parse_ast/lint 失败


def test_convert_retries_once_with_errors(client):
    c, _, app = client
    prov = SeqProvider([json.dumps(BROKEN, ensure_ascii=False), json.dumps(FIXTURE, ensure_ascii=False)])
    app.state.ai.provider = prov
    r = _convert(c)
    assert r.status_code == 202
    job = _wait(c, r.json()["id"])
    assert job["status"] == "done", job
    assert len(prov.calls) == 2
    assert prov.calls[1][0]["content"].startswith("[patent-studio:convert]")
    fix = prov.calls[1][-1]["content"]
    assert "校验失败" in fix and "上一轮输出" in fix and "权利要求" in fix
    # 第一轮提示词带了问卷与文件名
    first = prov.calls[0][-1]["content"]
    assert "scheduler paper.txt" in first and PAPER.strip()[:40] in first


def test_convert_draft_restore_promotes_staged_figures(client):
    """采纳转换草稿（恢复该 commit）时暂存的 cv-fig 自动晋级 figs/ 并入附图节——
    恢复别的版本不动，manifest 按 draft_sha 精确匹配后消费。"""
    c, projects, _ = client
    c.put("/api/projects/method/settings", json={"ai_enabled": True})
    # 先落一个非草稿版本，作为「恢复别的版本」的对照
    cur = c.get("/api/projects/method/content").json()["content"]
    c.put("/api/projects/method/content", json={"content": cur + "\n\n"})
    figs_dir = projects / "method" / "figs"
    before_figs = set(os.listdir(figs_dir))
    r = _convert(c, target="method", name="x.docx", data=_docx())
    assert r.status_code == 202, r.text
    job = _wait(c, r.json()["id"])
    assert job["status"] == "done", job
    draft_sha = job["result"]["draft_sha"]
    manifest = projects / "method" / "attachments" / "cv-fig-manifest.json"
    assert manifest.exists()

    # 恢复无关版本：暂存图留在 attachments，manifest 不动
    other = [v for v in c.get("/api/projects/method/history").json() if v["sha"] != draft_sha][-1]
    r = c.post(f"/api/projects/method/history/{other['sha']}/restore")
    assert r.status_code == 200, r.text
    assert r.json()["figures_promoted"] == [] and manifest.exists()
    assert (projects / "method" / "attachments" / "cv-fig1.png").exists()

    # 恢复草稿 commit：cv-fig1.png 晋级 figs/ 并编入附图节（fixture 声明 图1）
    r = c.post(f"/api/projects/method/history/{draft_sha}/restore")
    assert r.status_code == 200, r.text
    promoted = r.json()["figures_promoted"]
    assert promoted == ["cv-fig1.png"]
    assert (figs_dir / "cv-fig1.png").exists()
    assert not (projects / "method" / "attachments" / "cv-fig1.png").exists()
    assert not manifest.exists()
    doc = figures_md.read(str(projects / "method"))
    assert [f.file for f in doc.spec] == ["cv-fig1.png"]
    assert doc.abstract and doc.abstract.file == "cv-fig1.png"
    # 旧图留 figs/ 作素材不入书
    assert before_figs.issubset(set(os.listdir(figs_dir)))


def test_convert_fails_after_second_invalid(client):
    c, projects, app = client
    prov = SeqProvider(["not json at all", json.dumps(BROKEN, ensure_ascii=False)])
    app.state.ai.provider = prov
    r = _convert(c)
    job = _wait(c, r.json()["id"])
    assert job["status"] == "error" and job["error"]["kind"] == "ai_format" and len(prov.calls) == 2
    assert not (projects / "scheduler-paper" / "docs" / "claims.md").exists()


def test_order_figures_src_reorders_drops_and_fallbacks():
    """figures[].src（提取图序号，1 起）挑选重排提取图；未收录剔除；全无效/缺键/空表的回退语义。"""
    figs = [("png", b"a"), ("png", b"b"), ("png", b"c")]
    assert pipeline._order_figures(figs, [{"num": 1, "src": 3, "caption": "x"}]) == [("png", b"c")]
    assert pipeline._order_figures(figs, [{"src": 2}, {"src": 1}]) == [("png", b"b"), ("png", b"a")]
    assert pipeline._order_figures(figs, [{"src": 1}, {"src": 1}, {"src": 2}]) == [("png", b"a"), ("png", b"b")]
    assert pipeline._order_figures(figs, None) == figs  # 未输出 figures 键 → 原样全收
    assert pipeline._order_figures(figs, []) == []  # 模型判定全不收
    assert pipeline._order_figures(figs, [{"src": 99}, {"src": "x"}, {}]) == figs  # src 全无效 → 原序兑底


def test_order_figures_mmd_renders_gen_schematic():
    """figures[].mmd（自绘受限 mermaid）渲染为 PNG 收录，gen_src 收集图源；坏 mmd 跳过。"""
    mmd = "flowchart LR\n  A[输入] --> B[编码] --> C[输出]\n"
    gen = {}
    marks = []
    out = pipeline._order_figures(
        [], [{"num": 1, "mmd": mmd, "caption": "图1", "marks": {"101": "输入模块", "102": "编码模块"}}], gen, marks
    )
    assert len(out) == 1 and out[0][0] == "png" and out[0][1][:4] == b"\x89PNG"
    assert (gen.get("fig1.mmd") or "").strip() == mmd.strip()
    assert marks == [{"101": "输入模块", "102": "编码模块"}]
    # spec 图源走 figspec 确定性渲染，gen_src 落 .spec JSON
    spec = {
        "bands": [
            {"title": "输入", "blocks": [{"id": "A", "label": "输入模块 101"}]},
            {"title": "输出", "blocks": [{"id": "B", "label": "输出模块 102"}]},
        ],
        "edges": [{"from": "A", "to": "B"}],
    }
    gen2, marks2 = {}, []
    out2 = pipeline._order_figures(
        [], [{"num": 1, "spec": spec, "caption": "图1", "marks": {"101": "输入模块"}}], gen2, marks2
    )
    assert len(out2) == 1 and out2[0][0] == "png" and out2[0][1][:4] == b"\x89PNG"
    assert gen2.get("fig1.spec") and json.loads(gen2["fig1.spec"])["bands"]
    assert marks2 == [{"101": "输入模块"}]
    # 坏 mmd 渲染失败跳过；src/mmd 全无效 → 原序兑底
    figs = [("png", b"a")]
    assert pipeline._order_figures(figs, [{"num": 1, "mmd": "not a flowchart"}]) == figs
