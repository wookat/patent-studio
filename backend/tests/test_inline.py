# -*- coding: utf-8 -*-
"""W9 prep-kernel：行内格式 tokenizer（patent_studio.kernel.inline）单测 + 五种标记 DOCX / XML / 导入全链路。"""
import io
import os
import shutil

import pytest
from docx import Document
from lxml import etree

from patent_studio.kernel import cnipa_xml as cx
from patent_studio.kernel import critic, docx_import
from patent_studio.kernel.builder import DEFAULT_SUB_EXCEPT, PatentBuilder, build_patent_ex
from patent_studio.kernel.inline import MARKS, inline_spans, inline_strip, inline_tokens
from patent_studio.kernel.tracked import TrackedRuns

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
METHOD = os.path.join(ROOT, "examples", "method")

B, I, U, SUB, SUP = (frozenset({m}) for m in MARKS)
PLAIN = frozenset()

# 五种标记同段（供 DOCX + XML + 导入全链路）；导入器写下标时单字母数字优先 `_x`，当前文用 `_{2}` 则还原为 `_2`
FIVE = "参数**粗体**与*斜体*及<u>下划线</u>，变量x_t与e^2，嵌套**a_{2}**结束"
FIVE_CANON = FIVE.replace("a_{2}", "a_2")


def _toks(text, **kw):
    return inline_tokens(text, **kw)


def test_plain_and_empty():
    assert _toks("") == []
    assert _toks("纯文本，无标记。") == [(PLAIN, "纯文本，无标记。")]


def test_each_mark():
    assert _toks("**粗**") == [(B, "粗")]
    assert _toks("*斜*") == [(I, "斜")]
    assert _toks("<u>线</u>") == [(U, "线")]
    assert _toks("x_t") == [(PLAIN, "x"), (SUB, "t")]
    assert _toks("x_{t,j}y") == [(PLAIN, "x"), (SUB, "t,j"), (PLAIN, "y")]
    assert _toks("r_[k]") == [(PLAIN, "r"), (SUB, "[k]")]  # 与 builder 既有行为一致：方括号保留
    assert _toks("e^2") == [(PLAIN, "e"), (SUP, "2")]
    assert _toks("e^{n+1}") == [(PLAIN, "e"), (SUP, "n+1")]
    assert _toks("x_t,z_t") == [(PLAIN, "x"), (SUB, "t,z"), (SUB, "t")]  # 相邻下标不合并（零漂移）


def test_nested_marks():
    assert _toks("**a_{2}**") == [(B, "a"), (B | SUB, "2")]
    assert _toks("<u>**x**</u>") == [(B | U, "x")]
    assert _toks("*e^2*") == [(I, "e"), (I | SUP, "2")]
    assert _toks("***粗斜***") == [(B | I, "粗斜")]
    assert _toks("**粗*斜*粗**") == [(B, "粗"), (B | I, "斜"), (B, "粗")]


def test_star_boundaries_and_precedence():
    # `*` 邻空白不触发
    assert _toks("a * b * c") == [(PLAIN, "a * b * c")]
    assert _toks("* 斜*") == [(PLAIN, "* 斜*")]
    assert _toks("*斜 *") == [(PLAIN, "*斜 *")]
    # `**` 优先于 `*`
    assert _toks("**粗**") == [(B, "粗")]
    assert _toks("2**3") == [(PLAIN, "2**3")]  # 未闭合
    assert _toks("a*b*c") == [(PLAIN, "a"), (I, "b"), (PLAIN, "c")]
    # 不跨行
    assert _toks("**a\nb**") == [(PLAIN, "**a\nb**")]
    assert _toks("<u>a\nb</u>") == [(PLAIN, "<u>a\nb</u>")]
    # 未闭合 / 空
    assert _toks("<u></u>x") == [(PLAIN, "<u></u>x")]
    assert _toks("**") == [(PLAIN, "**")]


def test_sub_except_whitelist():
    text = "标识符tsf_guard与on_matured保持字面，但x_t是下标"
    assert _toks(text, sub_except=DEFAULT_SUB_EXCEPT) == [
        (PLAIN, "标识符tsf_guard与on_matured保持字面，但x"), (SUB, "t"), (PLAIN, "是下标")]
    assert _toks("tsf_guard") == [(PLAIN, "tsf"), (SUB, "guard")]
    # 白名单词在粗体里也保持字面
    assert _toks("**tsf_guard**", sub_except=DEFAULT_SUB_EXCEPT) == [(B, "tsf_guard")]


def test_greek_letter_scripts():
    # ε_θ / μ_θ / Σ_θ 类数学记号：希腊字母下标应转真下标而非露字面下划线
    # （ddpm-11 权项 Σ_θ 打印残渣实案）
    assert _toks("噪声ε_θ与均值μ_θ") == [
        (PLAIN, "噪声ε"), (SUB, "θ"), (PLAIN, "与均值μ"), (SUB, "θ")]
    assert _toks("协方差Σ_θ") == [(PLAIN, "协方差Σ"), (SUB, "θ")]
    assert _toks("指数x^α") == [(PLAIN, "指数x"), (SUP, "α")]


def test_criticmarkup_accept_then_tokenize():
    src = "误差{++**e_t**++}与{--旧--}{~~a~>*b*~~}"
    assert _toks(critic.accept(src)) == [(PLAIN, "误差"), (B, "e"), (B | SUB, "t"), (PLAIN, "与"), (I, "b")]
    assert _toks(critic.reject(src)) == [(PLAIN, "误差与旧a")]


def test_bare_script_token_greedy_needs_braces():
    """W9 domain-01：内核裸 `_x` 贪婪吞后续字母数字（这是既定语法，`x_t,c` 依赖它），
    故前端 / docx_import 序列化 H2O 只选 2 时必须写 `H_{2}O`；两种写法的解析结果在此钉住。"""
    assert _toks("水H_2O") == [(PLAIN, "水H"), (SUB, "2O")]
    assert _toks("水H_{2}O") == [(PLAIN, "水H"), (SUB, "2"), (PLAIN, "O")]
    assert _toks("m/s^2_2") == [(PLAIN, "m/s"), (SUP, "2"), (SUB, "2")]
    # docx_import 的序列化判据：后文会被多吃 → 花括号；否则裸 token
    assert docx_import._join_pieces([("t", "水H"), ("sub", "2"), ("t", "O")]) == "水H_{2}O"
    assert docx_import._join_pieces([("t", "x"), ("sub", "t"), ("t", "时刻")]) == "x_t时刻"
    assert docx_import._join_pieces([("t", "CO"), ("sub", "2"), ("t", "e")]) == "CO_{2}e"


def test_bold_italic_edge_uses_triple_star():
    """W9 ux-01：粗斜体紧贴粗体跨度边缘时 `**a*b***` 会被 `**` 懒匹配切成 b(a*b)+`*`（既定语法，不改内核），
    前端序列化改写为 `**a*****b***`（bi 分支）；此处钉住前端依赖的三种解析结果。"""
    assert _toks("**粗体*测试***") == [(B, "粗体*测试"), (PLAIN, "*")]  # 前端不得再产生这种写法
    assert _toks("**粗体*****测试***") == [(B, "粗体"), (B | I, "测试")]
    assert _toks("**粗体*****<u>测试</u>***") == [(B, "粗体"), (B | I | U, "测试")]
    assert _toks("***测试*****粗体**") == [(B | I, "测试"), (B, "粗体")]
    assert _toks("***测试****粗体*") == [(B | I, "测试"), (I, "粗体")]
    assert _toks("*a****b****c*") == [(I, "a"), (B | I, "b"), (I, "c")]
    assert _toks("**a*b*c**") == [(B, "a"), (B | I, "b"), (B, "c")]
    assert _toks("**a***b*") == [(B, "a"), (I, "b")]


def test_inline_strip_and_spans():
    assert inline_strip(FIVE) == "参数粗体与斜体及下划线，变量xt与e2，嵌套a2结束"
    assert inline_strip("tsf_guard", sub_except=DEFAULT_SUB_EXCEPT) == "tsf_guard"
    spans = inline_spans("**a_{2}**")
    assert spans == [{"text": "a", "marks": ["b"]}, {"text": "2", "marks": ["b", "sub"]}]
    assert all(set(s["marks"]) <= set(MARKS) for s in inline_spans(FIVE))


# ---------------- DOCX ----------------

def _runs(p):
    return [(r.text, bool(r.bold), bool(r.italic), bool(r.underline), bool(r.font.subscript),
             bool(r.font.superscript)) for r in p.runs]


def test_docx_runs_five_marks():
    b = PatentBuilder()
    p = b.para(FIVE)
    assert _runs(p) == [
        ("参数", False, False, False, False, False),
        ("粗体", True, False, False, False, False),
        ("与", False, False, False, False, False),
        ("斜体", False, True, False, False, False),
        ("及", False, False, False, False, False),
        ("下划线", False, False, True, False, False),
        ("，变量x", False, False, False, False, False),
        ("t", False, False, False, True, False),
        ("与e", False, False, False, False, False),
        ("2", False, False, False, False, True),
        ("，嵌套", False, False, False, False, False),
        ("a", True, False, False, False, False),
        ("2", True, False, False, True, False),
        ("结束", False, False, False, False, False),
    ]


def test_docx_runs_tracked_hook_keeps_marks():
    """run_hook（tracked.TrackedRuns）路径：{++**x**++} → w:ins 内的 run 仍带 bold。"""
    b = PatentBuilder(run_hook=TrackedRuns("审阅者", "2026-01-01T00:00:00Z"))
    p = b.para("前{++**粗**e^2++}后")
    ins = p._p.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ins")
    assert len(ins) == 1
    xml = etree.tostring(ins[0], encoding="unicode")
    assert "<w:b/>" in xml and 'w:vertAlign w:val="superscript"' in xml
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    assert [t.text for t in p._p.iter(W + "t")] == ["前", "粗", "e", "2", "后"]
    assert [t.text for t in ins[0].iter(W + "t")] == ["粗", "e", "2"]


# ---------------- CNIPA XML ----------------

def _xml(el):
    return etree.tostring(el, encoding="unicode")


def test_xml_inline_nested_tags():
    p = etree.Element("p")
    cx.inline_text(p, FIVE)
    assert _xml(p) == ("<p>参数<b>粗体</b>与<i>斜体</i>及<u>下划线</u>，变量x<sub>t</sub>与e<sup>2</sup>，"
                       "嵌套<b>a</b><b><sub>2</sub></b>结束</p>")
    p = etree.Element("p")
    cx.inline_text(p, "<u>***x***</u>y")
    assert _xml(p) == "<p><b><i><u>x</u></i></b>y</p>"


def test_xml_claim_refs_with_marks():
    ct = etree.Element("claim-text")
    cx.claim_text_with_refs(ct, "根据权利要求1所述的方法，其中**g_t**>0", 3, ())
    assert _xml(ct) == ('<claim-text>根据<claim-ref idref="c1">权利要求1</claim-ref>所述的方法，'
                        '其中<b>g</b><b><sub>t</sub></b>&gt;0</claim-text>')


def test_xml_five_marks_dtd_valid(tmp_path):
    ex = cx._Exporter({"sub_except": (), "style": None})
    root = ex.claims(["1.一种方法，其特征在于，" + FIVE, "2.根据权利要求1所述的方法，*其中*<u>**x_i**</u>。"])
    p = tmp_path / "c.xml"
    p.write_text(cx.serialize(root), encoding="utf-8")
    valid, errors = cx.validate_xml(str(p))
    assert valid, errors
    assert root.find(".//claim-text/b") is not None and root.find(".//claim-text/b/u/sub") is not None
    root, _ = ex.description("标题", [("h", "技术领域"), FIVE])
    p.write_text(cx.serialize(root), encoding="utf-8")
    valid, errors = cx.validate_xml(str(p))
    assert valid, errors


# ---------------- 导入（DOCX → Markdown）----------------

def test_import_marks_roundtrip_paragraph(tmp_path):
    """内核 build 的段 → import：五种标记逐字还原。"""
    b = PatentBuilder()
    b.para("摘要", bold=True)  # 书名段（仅供分书；导入器以字体/居中判定，这里走直接段读取）
    p = b.para(FIVE)
    para = docx_import._read_para(p._p, lambda om: "")
    assert docx_import._join_pieces(para.pieces, para.piece_marks) == FIVE_CANON
    # 嵌套顺序规范化为 b > i > u（由外向内）；上标多字符带花括号；空白留在记号外
    p = b.para("***粗斜***与<u>*x*</u>，e^{n+1}与 **两侧空白** 保留")
    para = docx_import._read_para(p._p, lambda om: "")
    assert docx_import._clean(docx_import._join_pieces(para.pieces, para.piece_marks)) == (
        "***粗斜***与*<u>x</u>*，e^{n+1}与 **两侧空白** 保留")


def test_import_native_word_marks(tmp_path):
    """Word 原生 run 格式（含跨软回车、首尾空白）→ 记号不跨行、空白留在记号外。"""
    d = Document()
    p = d.add_paragraph()
    r = p.add_run("粗体 ")
    r.bold = True
    r = p.add_run("续\n再")
    r.bold = True
    r = p.add_run("上")
    r.font.superscript = True
    r = p.add_run("下")
    r.italic = True
    r.underline = True
    para = docx_import._read_para(p._p, lambda om: "")
    assert docx_import._join_pieces(para.pieces, para.piece_marks) == "**粗体 续**\n**再**^{上}*<u>下</u>*"
    assert not any(k == "superscript" for k, _ in para.unsupported)


@pytest.mark.parametrize("name", ["method"])
def test_full_chain_example_with_marks(tmp_path, name):
    """真实案 + 五标记段：build → import → 标记段逐字还原；XML 导出结构（无 LibreOffice 渲染）含嵌套标签。"""
    from app.domain.projects import ProjectStore

    projects = tmp_path / "projects"
    projects.mkdir()
    shutil.copytree(os.path.join(ROOT, "examples", name), projects / name,
                    ignore=shutil.ignore_patterns("*.docx", ".studio", "*.bak-*"))
    store = ProjectStore(str(projects))
    proj = store.get(name)
    md_path = os.path.join(proj.dir, "content.md")
    md = open(md_path, encoding="utf-8").read()
    marker = "\n### 技术领域\n"
    assert marker in md
    md = md.replace(marker, marker + FIVE + "\n\n", 1)
    open(md_path, "w", encoding="utf-8").write(md)
    out = str(tmp_path / "m.docx")
    build_patent_ex(store.build_config(proj, out=out))
    imported, rep = docx_import.import_docx(open(out, "rb").read(), str(tmp_path / "figs"))
    assert FIVE_CANON in imported
    assert not any(i.code == "unsupported.superscript" for i in rep.items)
    doc = Document(io.BytesIO(open(out, "rb").read()))
    hit = [p for p in doc.paragraphs if p.text.endswith("嵌套a2结束")]
    assert hit and any(r.underline for r in hit[0].runs) and any(r.font.superscript for r in hit[0].runs)
    ex = cx._Exporter({"sub_except": tuple(store.load_settings(name)["sub_except"]), "style": None})
    from patent_studio.kernel.content import load_md
    title, _abstract, _claims, spec = load_md(md_path)
    root, _ = ex.description(title, spec)
    xml = cx.serialize(root)
    assert "<b>粗体</b>" in xml and "<i>斜体</i>" in xml and "<u>下划线</u>" in xml
    assert "<sub>t</sub>" in xml and "<sup>2</sup>" in xml and "<b><sub>2</sub></b>" in xml
    p = tmp_path / "d.xml"
    p.write_text(xml, encoding="utf-8")
    valid, errors = cx.validate_xml(str(p))
    assert valid, errors
